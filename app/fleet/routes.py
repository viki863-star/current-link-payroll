import sqlite3
import json
import base64
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from flask import (
    current_app, flash, redirect, render_template, request,
    send_file, url_for, session, Response
)
from werkzeug.security import generate_password_hash

from ..database import open_db
from ..routes import (
    _bulk_row_count,
    _bulk_row_values,
    _bulk_validation_errors,
    _ensure_maintenance_suppliers_table,
    _insert_staff_job_row,
    _login_required,
    _touch_admin_workspace,
    _upsert_maintenance_supplier,
)
from ..pdf_service import (
    BLUE, BLUE_DARK, BLUE_SOFT, LINE, MUTED, SOFT, TEXT,
    _draw_header, _draw_title, _fit_image_reader, _fit_text,
    format_currency, format_date_label,
    generate_fuel_report_pdf,
)
from . import fleet_bp


VEHICLE_TYPES = [
    "Tractor Unit",
    "Flat Bed Trailer",
    "Low Bed Trailer",
    "Container Trailer",
    "Curtain Side Trailer",
    "Box Trailer",
    "Tanker - Drinking Water",
    "Tanker - Non-Drinking Water",
    "Tanker - Drainage",
    "Tanker - Fuel/Oil",
    "Tanker - Chemical",
    "Pickup Truck (Single Cab)",
    "Pickup Truck (Double Cab)",
    "Box Truck",
    "Refrigerated Truck",
    "Flat Deck Truck",
    "Tipper Truck",
    "Crane",
    "Forklift",
    "Concrete Mixer",
    "Excavator",
    "Bulldozer",
    "Loader",
    "Roller",
    "Recovery / Tow Truck",
    "Garbage Truck",
    "Water Bowser",
    "Light Bus",
    "Heavy Bus",
    "Van / Minibus",
    "Car / SUV",
    "Other",
]
VEHICLE_CATEGORIES = ["Solo", "Head", "Trailer"]
VEHICLE_SUB_TYPES = ["Tractor", "Flat Bed 12M", "Flat Bed 24M", "Tanker Drinking", "Tanker Non-Drinking", "Tanker Drainage", "Box Truck", "Crane", "Forklift", "Other"]

# Divided registry sections — order matters (display order on /fleet/vehicles)
VEHICLE_GROUPS = [
    ("trailers", "Trailers", "🚜"),
    ("drinking", "Drinking Water Tankers", "💧"),
    ("non_drinking", "Non-Drinking Water Tankers", "🚰"),
    ("drainage", "Drainage Tankers", "🛢️"),
    ("cranes", "Cranes", "🏗️"),
    ("forklifts", "Forklifts", "🏗"),
    ("others", "Other Vehicles", "🚚"),
]


def _vehicle_group_key(v) -> str:
    """Which registry section a vehicle belongs to.

    Driven by vehicle_type / vehicle_sub_type / vehicle_category, so editing a
    vehicle's category/type moves it into the matching section automatically.
    Order matters: "non-drinking" contains the substring "drinking".
    """
    vt = (v.get("vehicle_type") or "").lower()
    st = (v.get("vehicle_sub_type") or "").lower()
    cat = (v.get("vehicle_category") or "").lower()
    if "non-drinking" in vt or "non drinking" in vt or "non-drinking" in st or "non drinking" in st:
        return "non_drinking"
    if "drainage" in vt or "drainage" in st:
        return "drainage"
    if "drinking" in vt or "drinking" in st:
        return "drinking"
    if "forklift" in vt or "forklift" in st:
        return "forklifts"
    if "crane" in vt or "crane" in st:
        return "cranes"
    if "trailer" in vt or cat == "trailer" or st in ("flat bed 12m", "flat bed 24m") or vt == "tractor unit" or st == "tractor":
        return "trailers"
    return "others"
LINK_TYPES = ["", "Flat Link", "Tanker Link"]
TANK_CAPACITIES = [0, 3000, 5000, 10000]
OWNERSHIP_TYPES = ["Standard", "Partnership"]
MAINTENANCE_CATEGORIES = ["Oil Change", "Tyre", "Engine", "Body", "Electrical", "Brakes", "AC", "Other"]

_MJ_LIST_COLS = """
    mj.id, mj.vehicle_id, mj.staff_id, mj.amount, mj.category, mj.description,
    mj.attachment_name, mj.attachment_type, mj.status, mj.admin_notes,
    mj.supplier_name, mj.supplier_trn, mj.supplier_bill_no, mj.tax_mode, mj.tax_amount,
    mj.created_at, mj.approved_at,
    CASE WHEN mj.attachment_data IS NOT NULL AND mj.attachment_data != '' THEN 1 ELSE 0 END AS has_attachment
"""


_fleet_tables_ensured = False

def ensure_fleet_tables():
    global _fleet_tables_ensured
    if _fleet_tables_ensured:
        return
    _fleet_tables_ensured = True
    db = open_db()
    db.execute("SELECT 1 FROM vehicles LIMIT 1")
    _migrate_vehicle_master(db)
    # Clean up blank staff entries
    db.execute("DELETE FROM field_staff WHERE staff_id IS NULL OR staff_id = ''")
    db.commit()
    # Drop FK constraints on PostgreSQL so staff can be deleted without losing data
    try:
        db.execute("ALTER TABLE maintenance_jobs DROP CONSTRAINT IF EXISTS maintenance_jobs_staff_id_fkey")
    except Exception:
        pass
    try:
        db.execute("ALTER TABLE maintenance_jobs DROP CONSTRAINT IF EXISTS fk_maintenance_jobs_staff_id")
    except Exception:
        pass
    try:
        db.execute("ALTER TABLE maintenance_jobs ALTER COLUMN staff_id DROP NOT NULL")
    except Exception:
        pass
    try:
        db.execute("ALTER TABLE maintenance_jobs ALTER COLUMN staff_id SET DEFAULT ''")
    except Exception:
        pass
    # Create vehicle_documents table
    id_col = "id INTEGER PRIMARY KEY AUTOINCREMENT" if db.backend == "sqlite" else "id SERIAL PRIMARY KEY"
    default_ts = "CURRENT_TIMESTAMP"
    db.execute(f"""
        CREATE TABLE IF NOT EXISTS vehicle_documents (
            {id_col},
            plate_no TEXT NOT NULL,
            doc_name TEXT NOT NULL,
            doc_type TEXT,
            doc_data TEXT,
            uploaded_at TEXT DEFAULT {default_ts},
            notes TEXT
        )
    """)
    # Migrate any old vehicle_documents records to unified documents table
    try:
        old_docs = db.execute("SELECT id, plate_no, doc_name, doc_type, doc_data, notes FROM vehicle_documents").fetchall()
        for od in old_docs:
            existing = db.execute("SELECT id FROM documents WHERE entity_type='vehicle' AND entity_id=? AND doc_name=? AND doc_category='Other'", (od["plate_no"], od["doc_name"])).fetchone()
            if not existing:
                db.execute(
                    "INSERT INTO documents (entity_type, entity_id, doc_name, doc_category, file_data, file_type, file_size, notes) VALUES ('vehicle',?,?,?,?,?,?,?)",
                    (od["plate_no"], od["doc_name"], "Other", od["doc_data"], od["doc_type"], len(od["doc_data"]) if od["doc_data"] else 0, od["notes"]),
                )
        db.commit()
    except Exception:
        pass
    real_type = "REAL"
    db.execute(f"""
        CREATE TABLE IF NOT EXISTS fuel_entries (
            {id_col},
            vehicle_plate TEXT NOT NULL,
            entry_date TEXT NOT NULL,
            gallons {real_type} NOT NULL,
            rate_per_gallon {real_type} NOT NULL,
            total_amount {real_type} NOT NULL,
            supplier_id INTEGER,
            supplier_name TEXT NOT NULL,
            notes TEXT,
            source_expense_id INTEGER,
            created_at TEXT DEFAULT {default_ts}
        )
    """)
    db.execute(f"""
        CREATE TABLE IF NOT EXISTS traffic_fines (
            {id_col},
            plate_no TEXT NOT NULL,
            fine_date TEXT NOT NULL,
            fine_type TEXT NOT NULL DEFAULT 'Traffic Fine',
            fine_amount {real_type} NOT NULL DEFAULT 0,
            fine_location TEXT,
            fine_description TEXT,
            fine_status TEXT NOT NULL DEFAULT 'Pending',
            attachment_data TEXT,
            attachment_type TEXT,
            attachment_name TEXT,
            notes TEXT,
            created_at TEXT DEFAULT {default_ts}
        )
    """)
    # Ghadeer cards (TAQA prepaid water-filling cards) + their ledger
    db.execute(f"""
        CREATE TABLE IF NOT EXISTS ghadeer_cards (
            {id_col},
            card_no TEXT NOT NULL UNIQUE,
            vehicle_plate TEXT,
            cardholder TEXT,
            status TEXT NOT NULL DEFAULT 'Active',
            notes TEXT,
            created_at TEXT DEFAULT {default_ts}
        )
    """)
    db.execute(f"""
        CREATE TABLE IF NOT EXISTS ghadeer_transactions (
            {id_col},
            card_id INTEGER NOT NULL,
            tx_type TEXT NOT NULL,
            amount {real_type} NOT NULL,
            volume_m3 {real_type},
            station TEXT,
            ref_no TEXT,
            tx_date TEXT NOT NULL,
            notes TEXT,
            created_at TEXT DEFAULT {default_ts}
        )
    """)
    db.execute("CREATE INDEX IF NOT EXISTS idx_gc_tx_card ON ghadeer_transactions (card_id)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_gc_tx_date ON ghadeer_transactions (tx_date)")
    # Ghadeer card photos: front/back scans stored base64 (same as documents.file_data)
    for _img_col in (
        "ALTER TABLE ghadeer_cards ADD COLUMN front_image TEXT",
        "ALTER TABLE ghadeer_cards ADD COLUMN front_image_type TEXT",
        "ALTER TABLE ghadeer_cards ADD COLUMN back_image TEXT",
        "ALTER TABLE ghadeer_cards ADD COLUMN back_image_type TEXT",
    ):
        try:
            db.execute(_img_col)
        except Exception:
            pass  # column already exists
    db.commit()
    # Fix existing fuel expenses — change earning_type from 'trip' to 'Fuel'
    try:
        db.execute(
            "UPDATE supplier_expenses SET earning_type = 'Fuel' WHERE category = 'Fuel' AND earning_type = 'trip'"
        )
        db.commit()
    except Exception:
        pass
    # Indexes to speed up list/profile queries
    for idx_sql in [
        "CREATE INDEX IF NOT EXISTS idx_mj_status_created ON maintenance_jobs (status, created_at)",
        "CREATE INDEX IF NOT EXISTS idx_mj_staff_created ON maintenance_jobs (staff_id, created_at)",
        "CREATE INDEX IF NOT EXISTS idx_mj_vehicle_status ON maintenance_jobs (vehicle_id, status)",
        "CREATE INDEX IF NOT EXISTS idx_msa_staff ON maintenance_staff_advances (staff_code)",
        "CREATE INDEX IF NOT EXISTS idx_mp_tech ON maintenance_papers (technician_code)",
    ]:
        try:
            db.execute(idx_sql)
        except Exception:
            pass
    try:
        db.commit()
    except Exception:
        pass


def _migrate_vehicle_master(db):
    """Copy vehicles from old vehicle_master table into vehicles table."""
    try:
        old = db.execute("SELECT id, vehicle_id, vehicle_no, vehicle_type, make_model, status, shift_mode, ownership_mode, source_type, source_party_code, source_asset_code, partner_party_code, partner_name, company_share_percent, partner_share_percent, notes FROM vehicle_master").fetchall()
    except Exception:
        return
    for v in old:
        existing = db.execute("SELECT plate_no FROM vehicles WHERE plate_no = ?", (v["vehicle_no"],)).fetchone()
        if existing:
            continue
        partner_percent = None
        try:
            partner_percent = float(v.get("partner_share_percent") or 0)
        except (ValueError, TypeError):
            pass
        try:
            db.execute(
                """INSERT INTO vehicles (plate_no, vehicle_type, model, ownership_type, partner_name, partner_percent, status, notes, created_at)
                   VALUES (?,?,?,?,?,?,?,?,COALESCE(?,CURRENT_TIMESTAMP))""",
                (v["vehicle_no"], v["vehicle_type"], v["make_model"],
                 v["ownership_mode"], v["partner_name"], partner_percent,
                 v["status"], v["notes"], v["created_at"]),
            )
        except Exception:
            pass
    db.commit()


def _vehicle_full(plate_no):
    db = open_db()
    v = db.execute("SELECT plate_no, vehicle_type, model, year, ownership_type, partner_name, partner_percent, status, notes FROM vehicles WHERE plate_no = ?", (plate_no,)).fetchone()
    if not v:
        vm = db.execute("SELECT id, vehicle_id, vehicle_no, vehicle_type, make_model, status, shift_mode, ownership_mode, source_type, source_party_code, source_asset_code, partner_party_code, partner_name, company_share_percent, partner_share_percent, notes FROM vehicle_master WHERE vehicle_no = ?", (plate_no,)).fetchone()
        if vm:
            v = {
                "plate_no": vm["vehicle_no"],
                "vehicle_type": vm["vehicle_type"],
                "model": vm["make_model"],
                "ownership_type": vm["ownership_mode"],
                "partner_name": vm["partner_name"],
                "partner_percent": vm.get("partner_share_percent"),
                "status": vm["status"],
                "notes": vm["notes"],
                "vehicle_id": vm["vehicle_id"],
            }
        else:
            return None
    driver = db.execute(
        """SELECT e.*, va.assigned_from FROM vehicle_assignments va
           JOIN employees e ON e.employee_id = va.driver_id
           WHERE va.vehicle_id = ? AND va.is_current = 1""",
        (plate_no,),
    ).fetchone()
    v["current_driver"] = driver
    job_count = db.execute(
        "SELECT COUNT(*) AS c FROM maintenance_jobs WHERE vehicle_id = ? AND status = 'approved'",
        (plate_no,),
    ).fetchone()["c"] or 0
    paper_count = db.execute(
        """SELECT COUNT(*) AS c FROM maintenance_papers mp
           JOIN vehicle_master vm ON vm.vehicle_id = mp.vehicle_id
           WHERE vm.vehicle_no = ? AND mp.review_status = 'Approved'""",
        (plate_no,),
    ).fetchone()["c"] or 0
    total_cost = db.execute(
        "SELECT COALESCE(SUM(amount),0) AS t FROM maintenance_jobs WHERE vehicle_id = ? AND status = 'approved'",
        (plate_no,),
    ).fetchone()["t"] or 0
    paper_cost = db.execute(
        """SELECT COALESCE(SUM(mp.total_amount),0) AS t FROM maintenance_papers mp
           JOIN vehicle_master vm ON vm.vehicle_id = mp.vehicle_id
           WHERE vm.vehicle_no = ? AND mp.review_status = 'Approved'""",
        (plate_no,),
    ).fetchone()["t"] or 0
    v["job_count"] = job_count + paper_count
    fuel_cost = 0
    try:
        fuel_cost = float(
            db.execute(
                "SELECT COALESCE(SUM(total_amount),0) AS t FROM fuel_entries WHERE vehicle_plate = ?",
                (plate_no,),
            ).fetchone()["t"] or 0
        )
    except Exception:
        pass
    parts_cost = 0
    try:
        parts_cost = float(
            db.execute(
                "SELECT COALESCE(SUM(net_amount),0) AS t FROM supplier_bills WHERE vehicle_plate = ?",
                (plate_no,),
            ).fetchone()["t"] or 0
        )
    except Exception:
        pass
    v["total_cost"] = float(total_cost) + float(paper_cost) + fuel_cost + parts_cost
    return v


def _all_employees_drivers():
    db = open_db()
    return db.execute(
        "SELECT employee_id, full_name FROM employees WHERE employee_type = 'Driver' AND status = 'Active' ORDER BY full_name"
    ).fetchall()


# ── Fleet Dashboard ─────────────────────────────────────────────

@fleet_bp.route("/fleet")
@_login_required("admin")
def fleet_dashboard():
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()

    vehicles = db.execute("SELECT plate_no, vehicle_type, model, year, ownership_type, partner_name, partner_percent, status, notes FROM vehicles WHERE plate_no NOT IN (SELECT linked_plate_no FROM vehicles WHERE linked_plate_no IS NOT NULL AND linked_plate_no != '') ORDER BY plate_no").fetchall()
    total = len(vehicles)
    active_v = sum(1 for v in vehicles if (v["status"] or "").lower() == "active")
    standard = sum(1 for v in vehicles if v["ownership_type"] == "Standard")
    partnership = sum(1 for v in vehicles if v["ownership_type"] == "Partnership")

    pending_jobs = db.execute(
        f"SELECT {_MJ_LIST_COLS}, COALESCE(v.plate_no, mj.vehicle_id) AS plate_no, v.vehicle_type, s.full_name AS staff_name FROM maintenance_jobs mj LEFT JOIN vehicles v ON v.plate_no = mj.vehicle_id JOIN field_staff s ON s.staff_id = mj.staff_id WHERE mj.status = 'pending' ORDER BY mj.created_at DESC"
    ).fetchall()

    pending_count = len(pending_jobs)

    total_maintenance_cost = db.execute(
        "SELECT COALESCE(SUM(amount),0) AS t FROM maintenance_jobs WHERE status = 'approved'"
    ).fetchone()["t"] or 0
    paper_cost = db.execute(
        "SELECT COALESCE(SUM(total_amount),0) AS t FROM maintenance_papers WHERE review_status = 'Approved'"
    ).fetchone()["t"] or 0
    total_maintenance_cost = float(total_maintenance_cost) + float(paper_cost)

    recent_jobs = db.execute(
        f"SELECT {_MJ_LIST_COLS}, COALESCE(v.plate_no, mj.vehicle_id) AS plate_no, v.vehicle_type, s.full_name AS staff_name FROM maintenance_jobs mj LEFT JOIN vehicles v ON v.plate_no = mj.vehicle_id JOIN field_staff s ON s.staff_id = mj.staff_id WHERE mj.status = 'approved' ORDER BY mj.created_at DESC LIMIT 10"
    ).fetchall()

    top_vehicles = db.execute(
        """SELECT COALESCE(v.plate_no, mj.vehicle_id) AS plate_no,
                  v.vehicle_type,
                  SUM(mj.amount) AS total_spent,
                  COUNT(mj.id) AS job_count
           FROM maintenance_jobs mj
           LEFT JOIN vehicles v ON v.plate_no = mj.vehicle_id
           WHERE mj.status = 'approved'
           GROUP BY COALESCE(v.plate_no, mj.vehicle_id), v.vehicle_type
           ORDER BY total_spent DESC
           LIMIT 10"""
    ).fetchall()

    total_pending_cost = db.execute(
        "SELECT COALESCE(SUM(amount),0) AS t FROM maintenance_jobs WHERE status='pending'"
    ).fetchone()["t"] or 0

    # Staff balances
    staff_balances = db.execute("""
        SELECT fs.staff_id, fs.full_name, fs.phone,
            COALESCE(adv.total_adv, 0) AS total_received,
            COALESCE(mj.total_jobs, 0) + COALESCE(mp.total_papers, 0) AS total_spent
        FROM field_staff fs
        LEFT JOIN (SELECT staff_code, SUM(amount) AS total_adv FROM maintenance_staff_advances GROUP BY staff_code) adv ON adv.staff_code = fs.staff_id
        LEFT JOIN (SELECT staff_id, SUM(COALESCE(staff_amount, amount - tax_amount)) AS total_jobs FROM maintenance_jobs WHERE status = 'approved' GROUP BY staff_id) mj ON mj.staff_id = fs.staff_id
        LEFT JOIN (SELECT technician_code, SUM(total_amount) AS total_papers FROM maintenance_papers WHERE review_status='Approved' GROUP BY technician_code) mp ON mp.technician_code = fs.staff_id
        WHERE fs.staff_id IS NOT NULL AND fs.staff_id != '' AND fs.staff_id != 'admin'
        ORDER BY fs.full_name
    """).fetchall()

    # Fuel stats
    from datetime import date as _date
    _today = _date.today()
    _cm = f"{_today.year:04d}-{_today.month:02d}"
    fuel_row = db.execute(
        "SELECT COUNT(*) AS entries, COALESCE(SUM(total_amount),0) AS total FROM fuel_entries WHERE LEFT(entry_date, 7) = %s",
        (_cm,),
    ).fetchone()
    fuel_entries_count = fuel_row["entries"] if fuel_row else 0
    fuel_total_amount = float(fuel_row["total"] if fuel_row else 0)

    # Chart data: monthly maintenance trend (last 6 months)
    _monthly_maint = db.execute("""
        SELECT TO_CHAR(created_at, 'YYYY-MM') AS ym,
               COALESCE(SUM(amount),0) AS total,
               COUNT(*) AS jobs
        FROM maintenance_jobs
        WHERE status = 'approved'
          AND created_at >= (CURRENT_DATE - INTERVAL '6 months')
        GROUP BY TO_CHAR(created_at, 'YYYY-MM')
        ORDER BY ym
    """).fetchall()
    chart_months = [r["ym"] for r in _monthly_maint]
    chart_costs = [float(r["total"]) for r in _monthly_maint]
    chart_job_counts = [r["jobs"] for r in _monthly_maint]

    # Chart data: vehicle status donut
    chart_vehicle_status = {
        "Active": sum(1 for v in vehicles if (v["status"] or "").lower() == "active"),
        "Standard": standard,
        "Partnership": partnership,
    }

    return render_template(
        "fleet/dashboard.html",
        vehicles=vehicles,
        total=total,
        active_count=active_v,
        standard_count=standard,
        partnership_count=partnership,
        pending_jobs=pending_jobs,
        pending_count=pending_count,
        total_maintenance_cost=total_maintenance_cost,
        recent_jobs=recent_jobs,
        top_vehicles=top_vehicles,
        total_pending_cost=float(total_pending_cost),
        staff_balances=staff_balances,
        fuel_entries_count=fuel_entries_count,
        fuel_total_amount=fuel_total_amount,
        chart_months=chart_months,
        chart_costs=chart_costs,
        chart_job_counts=chart_job_counts,
        chart_vehicle_status=chart_vehicle_status,
    )


# ── Vehicle List ────────────────────────────────────────────────

@fleet_bp.route("/fleet/vehicles")
@_login_required("admin")
def vehicle_list():
    try:
        _touch_admin_workspace("fleet")
        ensure_fleet_tables()
        db = open_db()

        q = request.args.get("q", "").strip()
        type_filter = request.args.get("type", "").strip()
        ownership_filter = request.args.get("ownership", "").strip()
        status_filter = request.args.get("status", "").strip()

        where = []
        params = []
        if q:
            where.append("(plate_no LIKE ? OR vehicle_type LIKE ? OR model LIKE ? OR partner_name LIKE ?)")
            like = f"%{q}%"
            params.extend([like, like, like, like])
        if type_filter:
            where.append("vehicle_type = ?")
            params.append(type_filter)
        if ownership_filter:
            where.append("ownership_type = ?")
            params.append(ownership_filter)
        if status_filter:
            where.append("v.status = ?")
            params.append(status_filter)

        # Default to Active only if no status filter
        if not status_filter:
            where.append("v.status = 'Active'")

        where_sql = " AND ".join(where) if where else "TRUE"

        vehicles = db.execute(
            f"""SELECT v.*, va.driver_id, e.full_name AS driver_name
                FROM vehicles v
                LEFT JOIN (
                    SELECT DISTINCT ON (vehicle_id) vehicle_id, driver_id
                    FROM vehicle_assignments
                    WHERE is_current = 1
                    ORDER BY vehicle_id, assigned_from DESC
                ) va ON va.vehicle_id = v.plate_no
                LEFT JOIN employees e ON e.employee_id = va.driver_id
                WHERE {where_sql}
                AND v.plate_no NOT IN (SELECT linked_plate_no FROM vehicles WHERE linked_plate_no IS NOT NULL AND linked_plate_no != '')
                ORDER BY v.ownership_type, v.plate_no""",
            params,

        ).fetchall()

        # Divided registry sections (Trailers / Drinking / Non-Drinking / Drainage / Crane / Forklift / Others)
        grouped = {key: [] for key, _, _ in VEHICLE_GROUPS}
        for v in vehicles:
            grouped[_vehicle_group_key(v)].append(v)
        groups = []
        group_counts = []
        for key, title, icon in VEHICLE_GROUPS:
            group_counts.append((key, title, icon, len(grouped[key])))
            if grouped[key]:
                groups.append((key, title, icon, grouped[key]))

        vehicle_types = [r[0] for r in db.execute("SELECT DISTINCT vehicle_type FROM vehicles ORDER BY vehicle_type").fetchall()]
        ownership_types = [r[0] for r in db.execute("SELECT DISTINCT ownership_type FROM vehicles ORDER BY ownership_type").fetchall()]
        active_count = db.execute("SELECT COUNT(*) FROM vehicles WHERE status = 'Active'").fetchone()[0]
        inactive_count = db.execute("SELECT COUNT(*) FROM vehicles WHERE status = 'Inactive'").fetchone()[0]
        stats = {
            "total": active_count + inactive_count,
            "active": active_count,
            "inactive": inactive_count,
            "heads": sum(1 for v in vehicles if (v.get("vehicle_category") or "Solo") == "Head"),
            "trailers": sum(1 for v in vehicles if (v.get("vehicle_category") or "Solo") == "Trailer"),
        }

        return render_template(
            "fleet/vehicle_list.html",
            vehicles=vehicles,
            groups=groups,
            group_counts=group_counts,
            stats=stats,
            q=q,
            type_filter=type_filter,
            ownership_filter=ownership_filter,
            status_filter=status_filter,
            vehicle_types=vehicle_types,
            ownership_types=ownership_types,
            VEHICLE_TYPES=VEHICLE_TYPES,
            OWNERSHIP_TYPES=OWNERSHIP_TYPES,
        )
    except Exception as e:
        current_app.logger.error("Fleet error: %s", e, exc_info=True)
        flash("An error occurred loading the fleet dashboard.", "error")
        return redirect(url_for("dashboard"))


# ═══════════════════════════════════════════════════════════════════
# GHADEER CARDS — TAQA Distribution prepaid water-filling cards
# Recharge = advance payment (credits balance), Fill = water drawn (debits).
# ═══════════════════════════════════════════════════════════════════

GHADEER_LOW_BALANCE = 1000.0


def _ghadeer_vehicle_options(db):
    """Active vehicles for the card-link dropdown: water tankers first."""
    t1, t2, t3 = "%tanker%", "%water bowser%", "%tanker%"
    tankers = db.execute(
        "SELECT plate_no, vehicle_type FROM vehicles WHERE status = 'Active' AND (lower(vehicle_type) LIKE ? OR lower(vehicle_type) LIKE ? OR lower(vehicle_sub_type) LIKE ?) ORDER BY plate_no",
        (t1, t2, t3),
    ).fetchall()
    others = db.execute(
        "SELECT plate_no, vehicle_type FROM vehicles WHERE status = 'Active' AND NOT (lower(vehicle_type) LIKE ? OR lower(vehicle_type) LIKE ? OR lower(vehicle_sub_type) LIKE ?) ORDER BY plate_no",
        (t1, t2, t3),
    ).fetchall()
    return tankers, others


def _ghadeer_card_or_none(db, card_id):
    return db.execute("SELECT * FROM ghadeer_cards WHERE id = ?", (card_id,)).fetchone()


@fleet_bp.route("/fleet/ghadeer-cards")
@_login_required("admin")
def ghadeer_cards():
    try:
        _touch_admin_workspace("fleet")
        ensure_fleet_tables()
        db = open_db()

        cards = db.execute(
            """SELECT c.id, c.card_no, c.vehicle_plate, c.cardholder, c.status, c.notes, c.created_at,
                      CASE WHEN c.front_image IS NOT NULL AND c.front_image != '' THEN 1 ELSE 0 END AS has_front,
                      CASE WHEN c.back_image IS NOT NULL AND c.back_image != '' THEN 1 ELSE 0 END AS has_back,
                      COALESCE(t.balance, 0) AS balance,
                      COALESCE(t.tx_count, 0) AS tx_count,
                      t.last_tx_date
               FROM ghadeer_cards c
               LEFT JOIN (
                   SELECT card_id, SUM(amount) AS balance, COUNT(*) AS tx_count, MAX(tx_date) AS last_tx_date
                   FROM ghadeer_transactions GROUP BY card_id
               ) t ON t.card_id = c.id
               ORDER BY CASE WHEN c.status = 'Active' THEN 0 ELSE 1 END, c.card_no"""
        ).fetchall()

        month_start = date.today().replace(day=1).isoformat()
        month_row = db.execute(
            """SELECT COALESCE(SUM(CASE WHEN amount > 0 THEN amount END), 0) AS in_amt,
                      COALESCE(SUM(CASE WHEN amount < 0 THEN -amount END), 0) AS out_amt
               FROM ghadeer_transactions WHERE tx_date >= ?""",
            (month_start,),
        ).fetchone()

        tanker_vehicles, other_vehicles = _ghadeer_vehicle_options(db)

        return render_template(
            "fleet/ghadeer_cards.html",
            cards=cards,
            month_in=float(month_row["in_amt"] or 0),
            month_out=float(month_row["out_amt"] or 0),
            total_balance=sum(float(c["balance"] or 0) for c in cards),
            active_count=sum(1 for c in cards if c["status"] == "Active"),
            low_balance=GHADEER_LOW_BALANCE,
            tanker_vehicles=tanker_vehicles,
            other_vehicles=other_vehicles,
            today=date.today().isoformat(),
        )
    except Exception as e:
        current_app.logger.error("Ghadeer cards error: %s", e, exc_info=True)
        flash("An error occurred loading Ghadeer cards.", "error")
        return redirect(url_for("fleet.fleet_dashboard"))


@fleet_bp.route("/fleet/ghadeer-cards/add", methods=["POST"])
@_login_required("admin")
def ghadeer_card_add():
    ensure_fleet_tables()
    db = open_db()
    card_no = (request.form.get("card_no") or "").strip()
    if not card_no:
        flash("Ghadeer card number is required.", "error")
        return redirect(url_for("fleet.ghadeer_cards"))
    if db.execute("SELECT id FROM ghadeer_cards WHERE lower(card_no) = lower(?)", (card_no,)).fetchone():
        flash(f"Ghadeer card '{card_no}' already exists.", "error")
        return redirect(url_for("fleet.ghadeer_cards"))
    try:
        opening = abs(float(request.form.get("opening_balance") or 0))
    except (TypeError, ValueError):
        opening = 0.0
    params = (
        card_no,
        (request.form.get("vehicle_plate") or "").strip() or None,
        (request.form.get("cardholder") or "").strip() or None,
        (request.form.get("notes") or "").strip() or None,
    )
    if db.backend == "postgres":
        card_id = db.execute(
            "INSERT INTO ghadeer_cards (card_no, vehicle_plate, cardholder, status, notes) VALUES (?, ?, ?, 'Active', ?) RETURNING id",
            params,
        ).fetchone()[0]
    else:
        db.execute(
            "INSERT INTO ghadeer_cards (card_no, vehicle_plate, cardholder, status, notes) VALUES (?, ?, ?, 'Active', ?)",
            params,
        )
        card_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
    if opening > 0:
        db.execute(
            "INSERT INTO ghadeer_transactions (card_id, tx_type, amount, tx_date, notes) VALUES (?, 'Recharge', ?, ?, 'Opening balance')",
            (card_id, opening, (request.form.get("tx_date") or "").strip() or date.today().isoformat()),
        )
    db.commit()
    flash(f"Ghadeer card {card_no} added.", "success")
    return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))


@fleet_bp.route("/fleet/ghadeer-cards/<int:card_id>")
@_login_required("admin")
def ghadeer_card_detail(card_id):
    try:
        _touch_admin_workspace("fleet")
        ensure_fleet_tables()
        db = open_db()
        card = _ghadeer_card_or_none(db, card_id)
        if not card:
            flash("Ghadeer card not found.", "error")
            return redirect(url_for("fleet.ghadeer_cards"))
        card = dict(card)
        has_front = bool(card.get("front_image"))
        has_back = bool(card.get("back_image"))
        for _k in ("front_image", "front_image_type", "back_image", "back_image_type"):
            card.pop(_k, None)
        rows = db.execute(
            "SELECT * FROM ghadeer_transactions WHERE card_id = ? ORDER BY tx_date DESC, id DESC",
            (card_id,),
        ).fetchall()
        # running balance: walk ascending, display newest first
        running = 0.0
        ledger = []
        for t in reversed(rows):
            running += float(t["amount"] or 0)
            item = dict(t)
            item["balance"] = running
            ledger.append(item)
        ledger.reverse()
        total_recharged = sum(float(t["amount"] or 0) for t in rows if float(t["amount"] or 0) > 0)
        total_consumed = -sum(float(t["amount"] or 0) for t in rows if float(t["amount"] or 0) < 0)
        linked_vehicle = None
        if card["vehicle_plate"]:
            linked_vehicle = db.execute("SELECT plate_no FROM vehicles WHERE plate_no = ?", (card["vehicle_plate"],)).fetchone()
        tanker_vehicles, other_vehicles = _ghadeer_vehicle_options(db)
        return render_template(
            "fleet/ghadeer_card_detail.html",
            card=card,
            has_front=has_front,
            has_back=has_back,
            ledger=ledger,
            balance=running,
            total_recharged=total_recharged,
            total_consumed=total_consumed,
            low_balance=GHADEER_LOW_BALANCE,
            linked_vehicle=linked_vehicle,
            tanker_vehicles=tanker_vehicles,
            other_vehicles=other_vehicles,
            today=date.today().isoformat(),
        )
    except Exception as e:
        current_app.logger.error("Ghadeer card detail error: %s", e, exc_info=True)
        flash("An error occurred loading the Ghadeer card.", "error")
        return redirect(url_for("fleet.ghadeer_cards"))


@fleet_bp.route("/fleet/ghadeer-cards/<int:card_id>/tx", methods=["POST"])
@_login_required("admin")
def ghadeer_card_tx(card_id):
    ensure_fleet_tables()
    db = open_db()
    card = _ghadeer_card_or_none(db, card_id)
    if not card:
        flash("Ghadeer card not found.", "error")
        return redirect(url_for("fleet.ghadeer_cards"))
    tx_type = (request.form.get("tx_type") or "").strip()
    if tx_type not in ("Recharge", "Fill", "Adjustment"):
        flash("Invalid transaction type.", "error")
        return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))
    try:
        amount = abs(float(request.form.get("amount") or 0))
    except (TypeError, ValueError):
        amount = 0.0
    if amount <= 0:
        flash("Amount must be greater than zero.", "error")
        return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))
    is_debit = tx_type == "Fill" or (tx_type == "Adjustment" and request.form.get("direction") == "debit")
    try:
        volume = float(request.form.get("volume_m3") or 0) or None
    except (TypeError, ValueError):
        volume = None
    db.execute(
        """INSERT INTO ghadeer_transactions (card_id, tx_type, amount, volume_m3, station, ref_no, tx_date, notes)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            card_id,
            tx_type,
            -amount if is_debit else amount,
            volume,
            (request.form.get("station") or "").strip() or None,
            (request.form.get("ref_no") or "").strip() or None,
            (request.form.get("tx_date") or "").strip() or date.today().isoformat(),
            (request.form.get("notes") or "").strip() or None,
        ),
    )
    db.commit()
    flash(f"{tx_type} of AED {amount:,.2f} recorded on card {card['card_no']}.", "success")
    return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))


@fleet_bp.route("/fleet/ghadeer-cards/<int:card_id>/tx/<int:tx_id>/delete", methods=["POST"])
@_login_required("admin")
def ghadeer_card_tx_delete(card_id, tx_id):
    ensure_fleet_tables()
    db = open_db()
    tx = db.execute(
        "SELECT id, tx_type, amount FROM ghadeer_transactions WHERE id = ? AND card_id = ?",
        (tx_id, card_id),
    ).fetchone()
    if not tx:
        flash("Transaction not found.", "error")
    else:
        db.execute("DELETE FROM ghadeer_transactions WHERE id = ?", (tx_id,))
        db.commit()
        flash(f"{tx['tx_type']} entry deleted.", "success")
    return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))


@fleet_bp.route("/fleet/ghadeer-cards/<int:card_id>/edit", methods=["POST"])
@_login_required("admin")
def ghadeer_card_edit(card_id):
    ensure_fleet_tables()
    db = open_db()
    card = _ghadeer_card_or_none(db, card_id)
    if not card:
        flash("Ghadeer card not found.", "error")
        return redirect(url_for("fleet.ghadeer_cards"))
    card_no = (request.form.get("card_no") or "").strip()
    if not card_no:
        flash("Card number is required.", "error")
        return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))
    if db.execute("SELECT id FROM ghadeer_cards WHERE lower(card_no) = lower(?) AND id != ?", (card_no, card_id)).fetchone():
        flash(f"Card number '{card_no}' is already used by another card.", "error")
        return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))
    status = (request.form.get("status") or "Active").strip()
    if status not in ("Active", "Inactive", "Lost"):
        status = "Active"
    db.execute(
        "UPDATE ghadeer_cards SET card_no = ?, vehicle_plate = ?, cardholder = ?, status = ?, notes = ? WHERE id = ?",
        (
            card_no,
            (request.form.get("vehicle_plate") or "").strip() or None,
            (request.form.get("cardholder") or "").strip() or None,
            status,
            (request.form.get("notes") or "").strip() or None,
            card_id,
        ),
    )
    db.commit()
    flash("Ghadeer card updated.", "success")
    return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))


@fleet_bp.route("/fleet/ghadeer-cards/<int:card_id>/delete", methods=["POST"])
@_login_required("admin")
def ghadeer_card_delete(card_id):
    ensure_fleet_tables()
    db = open_db()
    card = _ghadeer_card_or_none(db, card_id)
    if not card:
        flash("Ghadeer card not found.", "error")
        return redirect(url_for("fleet.ghadeer_cards"))
    tx_count = db.execute("SELECT COUNT(*) FROM ghadeer_transactions WHERE card_id = ?", (card_id,)).fetchone()[0]
    if tx_count:
        flash("This card has transactions — set its status to Inactive instead of deleting.", "error")
        return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))
    db.execute("DELETE FROM ghadeer_cards WHERE id = ?", (card_id,))
    db.commit()
    flash(f"Ghadeer card {card['card_no']} deleted.", "success")
    return redirect(url_for("fleet.ghadeer_cards"))


# ── Ghadeer card photos (front/back) + PDF ─────────────────────────────────

_GHADEER_IMAGE_TYPES = {
    "image/jpeg": b"\xff\xd8\xff",
    "image/png": b"\x89PNG",
    "image/gif": b"GIF8",
    "image/webp": b"RIFF",
}
_GHADEER_IMAGE_MAX = 8 * 1024 * 1024  # 8 MB per photo


def _ghadeer_validate_image(data: bytes, mimetype: str) -> str:
    """Return the normalized mimetype or raise ValueError with a user message."""
    mt = (mimetype or "").lower()
    if mt not in _GHADEER_IMAGE_TYPES:
        raise ValueError("Only JPG, PNG, WEBP or GIF images are allowed.")
    if not data or len(data) > _GHADEER_IMAGE_MAX:
        raise ValueError("Image is empty or too large (max 8 MB).")
    if not data.startswith(_GHADEER_IMAGE_TYPES[mt]):
        raise ValueError("File does not look like a valid image.")
    if mt == "image/webp" and data[8:12] != b"WEBP":
        raise ValueError("File does not look like a valid image.")
    return mt


@fleet_bp.route("/fleet/ghadeer-cards/<int:card_id>/image/<side>", methods=["POST"])
@_login_required("admin")
def ghadeer_card_image(card_id, side):
    ensure_fleet_tables()
    db = open_db()
    card = _ghadeer_card_or_none(db, card_id)
    if not card:
        flash("Ghadeer card not found.", "error")
        return redirect(url_for("fleet.ghadeer_cards"))
    if side not in ("front", "back"):
        flash("Invalid photo side.", "error")
        return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))

    if (request.form.get("action") or "") == "delete":
        db.execute(
            f"UPDATE ghadeer_cards SET {side}_image = NULL, {side}_image_type = NULL WHERE id = ?",
            (card_id,),
        )
        db.commit()
        flash(f"{side.title()} side photo removed.", "success")
        return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))

    f = request.files.get("image")
    if not f or not f.filename:
        flash("Choose an image file first.", "error")
        return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))
    try:
        data = f.read()
        mt = _ghadeer_validate_image(data, f.mimetype)
    except ValueError as e:
        flash(str(e), "error")
        return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))
    db.execute(
        f"UPDATE ghadeer_cards SET {side}_image = ?, {side}_image_type = ? WHERE id = ?",
        (base64.b64encode(data).decode("ascii"), mt, card_id),
    )
    db.commit()
    flash(f"{side.title()} side photo saved.", "success")
    return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))


@fleet_bp.route("/fleet/ghadeer-cards/<int:card_id>/image/<side>")
@_login_required("admin")
def ghadeer_card_image_view(card_id, side):
    if side not in ("front", "back"):
        return Response(status=404)
    ensure_fleet_tables()
    db = open_db()
    row = db.execute(
        f"SELECT {side}_image AS img, {side}_image_type AS mt FROM ghadeer_cards WHERE id = ?",
        (card_id,),
    ).fetchone()
    if not row or not row["img"]:
        return Response(status=404)
    try:
        data = base64.b64decode(row["img"])
    except Exception:
        return Response(status=404)
    return send_file(
        BytesIO(data),
        mimetype=row["mt"] or "image/jpeg",
        conditional=True,
        max_age=3600,
    )


def _ghadeer_pdf_bytes(card, ledger, balance, total_recharged, total_consumed, company_profile) -> bytes:
    """One-page-style Ghadeer card statement (multi-page when the ledger is long)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as rl_canvas

    page_w, page_h = A4
    buf = BytesIO()
    pdf = rl_canvas.Canvas(buf, pagesize=A4)

    _draw_header(pdf, "", company_profile)
    subtitle = " · ".join(
        x for x in [
            str(card.get("card_no") or ""),
            card.get("vehicle_plate") or "No vehicle",
            card.get("cardholder") or "",
            date.today().isoformat(),
        ] if x
    )
    _draw_title(pdf, "Ghadeer Card Statement", subtitle)

    left = 16 * mm
    width = 178 * mm

    # ── summary strip ──
    y = page_h - 76 * mm
    summary_h = 26 * mm
    pdf.setFillColor(colors.white)
    pdf.setStrokeColor(LINE)
    pdf.roundRect(left, y - summary_h, width, summary_h, 3 * mm, fill=1, stroke=1)
    cells = [
        ("CURRENT BALANCE", "%.2f" % balance, True),
        ("RECHARGED", "+ %.2f" % total_recharged, False),
        ("FILLED", "- %.2f" % total_consumed, False),
        ("ENTRIES", str(len(ledger)), False),
    ]
    cw = width / 4
    for i, (lbl, val, is_balance) in enumerate(cells):
        cx = left + i * cw
        if i:
            pdf.setStrokeColor(LINE)
            pdf.line(cx, y - summary_h + 3 * mm, cx, y - 3 * mm)
        pdf.setFillColor(MUTED)
        pdf.setFont("Helvetica-Bold", 7)
        pdf.drawCentredString(cx + cw / 2, y - 8 * mm, lbl)
        if is_balance and balance < 0:
            pdf.setFillColor(colors.HexColor("#C92A2A"))
        elif is_balance:
            pdf.setFillColor(colors.HexColor("#2B8A3E"))
        else:
            pdf.setFillColor(BLUE_DARK)
        vtxt, vsize = _fit_text(pdf, val, "Helvetica-Bold", 14, cw - 6 * mm, min_size=8)
        pdf.setFont("Helvetica-Bold", vsize)
        pdf.drawCentredString(cx + cw / 2, y - 18 * mm, vtxt)
    y -= summary_h + 6 * mm

    # ── card photos ──
    photo_h = 52 * mm
    gap = 6 * mm
    box_w = (width - gap) / 2
    for i, (side, label) in enumerate((("front", "FRONT SIDE"), ("back", "BACK SIDE"))):
        bx = left + i * (box_w + gap)
        pdf.setFillColor(SOFT)
        pdf.setStrokeColor(LINE)
        pdf.roundRect(bx, y - photo_h, box_w, photo_h, 2.5 * mm, fill=1, stroke=1)
        pdf.setFillColor(BLUE_SOFT)
        pdf.roundRect(bx + 1.5 * mm, y - 7.5 * mm, box_w - 3 * mm, 6 * mm, 1.5 * mm, fill=1, stroke=0)
        pdf.setFillColor(BLUE_DARK)
        pdf.setFont("Helvetica-Bold", 7.5)
        pdf.drawCentredString(bx + box_w / 2, y - 5.8 * mm, label)
        drew = False
        img_b64 = card.get(side + "_image")
        if img_b64:
            try:
                raw = base64.b64decode(img_b64)
                reader = _fit_image_reader(raw, 1000)
                iw, ih = reader.getSize()
                avail_w = box_w - 6 * mm
                avail_h = photo_h - 14 * mm
                scale = min(avail_w / iw, avail_h / ih)
                dw, dh = iw * scale, ih * scale
                pdf.drawImage(
                    reader,
                    bx + (box_w - dw) / 2,
                    y - photo_h + 3 * mm + (avail_h - dh) / 2,
                    width=dw, height=dh,
                    preserveAspectRatio=True, mask="auto",
                )
                drew = True
            except Exception:
                drew = False
        if not drew:
            pdf.setFillColor(MUTED)
            pdf.setFont("Helvetica", 8)
            pdf.drawCentredString(bx + box_w / 2, y - photo_h / 2 - 2 * mm, "No photo uploaded")
    y -= photo_h + 8 * mm

    # ── transactions ──
    cols = [
        ("Date", 16, 21, "l"),
        ("Type", 37, 24, "l"),
        ("Description", 61, 63, "l"),
        ("Credit", 124, 23, "r"),
        ("Debit", 147, 21, "r"),
        ("Balance", 168, 26, "r"),
    ]
    row_h = 6.2 * mm
    hdr_h = 6.5 * mm
    page_no = 1
    C_GREEN = colors.HexColor("#2B8A3E")
    C_ORANGE = colors.HexColor("#E67700")
    C_RED = colors.HexColor("#C92A2A")
    type_colors = {
        "Recharge": C_GREEN,
        "Fill": colors.HexColor("#1971C2"),
        "Adjustment": C_ORANGE,
    }

    def footer():
        pdf.setFillColor(MUTED)
        pdf.setFont("Helvetica", 6.5)
        pdf.drawCentredString(
            page_w / 2, 8 * mm,
            "Generated %s · Current Link ERP · Page %d" % (date.today().isoformat(), page_no),
        )

    def table_header(yy):
        pdf.setFillColor(BLUE)
        pdf.roundRect(left, yy - hdr_h, width, hdr_h, 1.5 * mm, fill=1, stroke=0)
        pdf.setFillColor(colors.white)
        pdf.setFont("Helvetica-Bold", 7.4)
        for lbl, x_mm, wd, al in cols:
            if al == "r":
                pdf.drawRightString((x_mm + wd) * mm - 2.2 * mm, yy - 4.6 * mm, lbl)
            else:
                pdf.drawString(x_mm * mm + 2.2 * mm, yy - 4.6 * mm, lbl)
        return yy - hdr_h

    def new_page():
        nonlocal page_no, y
        footer()
        pdf.showPage()
        page_no += 1
        pdf.setFillColor(BLUE_DARK)
        pdf.setFont("Helvetica-Bold", 9.5)
        pdf.drawString(left, page_h - 20 * mm, "Ghadeer Card %s — transactions (continued)" % card.get("card_no", ""))
        y = table_header(page_h - 26 * mm)

    pdf.setFillColor(BLUE_DARK)
    pdf.setFont("Helvetica-Bold", 9.5)
    pdf.drawString(left, y - 1 * mm, "TRANSACTIONS")
    y -= 6 * mm
    y = table_header(y)

    def num(v):
        try:
            return "%g" % float(v)
        except (TypeError, ValueError):
            return ""

    if not ledger:
        pdf.setFillColor(MUTED)
        pdf.setFont("Helvetica", 8.5)
        pdf.drawCentredString(page_w / 2, y - 10 * mm, "No transactions yet.")
        y -= 16 * mm
    for idx, item in enumerate(ledger):
        if y - row_h < 16 * mm:
            new_page()
        if idx % 2 == 1:
            pdf.setFillColor(SOFT)
            pdf.roundRect(left, y - row_h, width, row_h, 1.2 * mm, fill=1, stroke=0)
        amt = float(item.get("amount") or 0)
        desc_bits = []
        if item.get("volume_m3"):
            desc_bits.append("%s m³" % num(item.get("volume_m3")))
        if item.get("station"):
            desc_bits.append(str(item["station"]))
        if item.get("ref_no"):
            desc_bits.append("#%s" % item["ref_no"])
        if item.get("notes"):
            desc_bits.append(str(item["notes"]))
        values = [
            str(item.get("tx_date") or ""),
            str(item.get("tx_type") or ""),
            " · ".join(desc_bits) or "-",
            ("%.2f" % amt) if amt > 0 else "",
            ("%.2f" % -amt) if amt < 0 else "",
            "%.2f" % float(item.get("balance") or 0),
        ]
        baseline = y - 4.3 * mm
        bal = float(item.get("balance") or 0)
        # single pass with an explicit colour per cell (no leftover fill state)
        styles = [
            ("Helvetica", 7.0, TEXT),
            ("Helvetica-Bold", 7.0, type_colors.get(item.get("tx_type"), TEXT)),
            ("Helvetica", 6.8, TEXT),
            ("Helvetica-Bold", 7.0, C_GREEN if amt > 0 else TEXT),
            ("Helvetica-Bold", 7.0, C_ORANGE if amt < 0 else TEXT),
            ("Helvetica-Bold", 7.0, C_RED if bal < 0 else BLUE_DARK),
        ]
        for (val, (_, x_mm, wd, al), (fname, fsize, col)) in zip(values, cols, styles):
            cell_x = x_mm * mm
            cell_w = wd * mm
            txt, sz = _fit_text(pdf, str(val), fname, fsize, cell_w - 4 * mm, min_size=5.4)
            pdf.setFont(fname, sz)
            pdf.setFillColor(col)
            if al == "r":
                pdf.drawRightString(cell_x + cell_w - 2 * mm, baseline, txt)
            else:
                pdf.drawString(cell_x + 2 * mm, baseline, txt)
        y -= row_h

    # ── totals ──
    if ledger:
        if y - 9 * mm < 16 * mm:
            new_page()
        y -= 2 * mm
        pdf.setFillColor(BLUE_SOFT)
        pdf.setStrokeColor(LINE)
        pdf.roundRect(left, y - 8 * mm, width, 8.5 * mm, 1.5 * mm, fill=1, stroke=1)
        pdf.setFillColor(BLUE_DARK)
        pdf.setFont("Helvetica-Bold", 8)
        pdf.drawString(left + 3 * mm, y - 5.6 * mm, "TOTALS")
        pdf.setFillColor(colors.HexColor("#2B8A3E"))
        pdf.drawString(126 * mm, y - 5.6 * mm, "+ %.2f" % total_recharged)
        pdf.setFillColor(colors.HexColor("#E67700"))
        pdf.drawString(150 * mm, y - 5.6 * mm, "- %.2f" % total_consumed)
        pdf.setFillColor(BLUE_DARK)
        pdf.setFont("Helvetica-Bold", 8.5)
        pdf.drawRightString(190 * mm, y - 5.6 * mm, "Balance: %.2f" % balance)
        y -= 12 * mm

    pdf.setFillColor(MUTED)
    pdf.setFont("Helvetica", 6.8)
    pdf.drawString(left, y - 2 * mm, "Card: %s" % card.get("card_no", ""))
    if card.get("notes"):
        pdf.drawString(left + 55 * mm, y - 2 * mm, "Notes: %s" % str(card["notes"])[:90])
    footer()
    pdf.save()
    return buf.getvalue()


@fleet_bp.route("/fleet/ghadeer-cards/<int:card_id>/pdf")
@_login_required("admin")
def ghadeer_card_pdf(card_id):
    try:
        ensure_fleet_tables()
        db = open_db()
        card = _ghadeer_card_or_none(db, card_id)
        if not card:
            flash("Ghadeer card not found.", "error")
            return redirect(url_for("fleet.ghadeer_cards"))
        card = dict(card)
        rows = db.execute(
            "SELECT * FROM ghadeer_transactions WHERE card_id = ? ORDER BY tx_date, id",
            (card_id,),
        ).fetchall()
        running = 0.0
        ledger = []
        for t in rows:
            running += float(t["amount"] or 0)
            item = dict(t)
            item["balance"] = running
            ledger.append(item)
        total_recharged = sum(float(t["amount"] or 0) for t in rows if float(t["amount"] or 0) > 0)
        total_consumed = -sum(float(t["amount"] or 0) for t in rows if float(t["amount"] or 0) < 0)

        try:
            cp = db.execute(
                "SELECT company_name, address, phone_number, email, trn_no, logo_data, logo_type "
                "FROM company_profile LIMIT 1"
            ).fetchone()
            cp = dict(cp) if cp else None
        except Exception:
            cp = None

        pdf_bytes = _ghadeer_pdf_bytes(
            card, ledger, running, total_recharged, total_consumed, cp
        )
        fname = "Ghadeer_Card_%s.pdf" % str(card.get("card_no") or card_id).replace(" ", "_").replace("/", "-")
        return send_file(
            BytesIO(pdf_bytes),
            mimetype="application/pdf",
            as_attachment=True,
            download_name=fname,
        )
    except Exception as e:
        current_app.logger.error("Ghadeer PDF error: %s", e, exc_info=True)
        flash("Could not generate the PDF.", "error")
        return redirect(url_for("fleet.ghadeer_card_detail", card_id=card_id))


@fleet_bp.route("/fleet/vehicles/download/excel")
@_login_required("admin")
def vehicle_list_excel():
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()

    q = request.args.get("q", "").strip()
    type_filter = request.args.get("type", "").strip()
    ownership_filter = request.args.get("ownership", "").strip()
    status_filter = request.args.get("status", "").strip()

    where = []
    params = []
    if q:
        where.append("(v.plate_no LIKE ? OR v.vehicle_type LIKE ? OR v.model LIKE ? OR v.partner_name LIKE ?)")
        like = f"%{q}%"
        params.extend([like, like, like, like])
    if type_filter:
        where.append("v.vehicle_type = ?")
        params.append(type_filter)
    if ownership_filter:
        where.append("v.ownership_type = ?")
        params.append(ownership_filter)
    if status_filter:
        where.append("v.status = ?")
        params.append(status_filter)

    where_sql = " AND ".join(where) if where else "TRUE"

    vehicles = db.execute(
        f"""SELECT v.*, va.driver_id, e.full_name AS driver_name,
                   lv.vehicle_type AS linked_vehicle_type, lv.model AS linked_vehicle_model
            FROM vehicles v
            LEFT JOIN vehicle_assignments va ON va.vehicle_id = v.plate_no AND va.is_current = 1
            LEFT JOIN employees e ON e.employee_id = va.driver_id
            LEFT JOIN vehicles lv ON lv.plate_no = v.linked_plate_no
            WHERE {where_sql}
            ORDER BY v.vehicle_category, v.plate_no""",
        params,
    ).fetchall()

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    ws.title = "Vehicles"

    hf = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    hfill = PatternFill("solid", fgColor="1a3a5c")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="d8e4f5")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)

    heads = ["#", "Plate No", "Category", "Type", "Sub Type", "Model", "Year", "Length", "Tank Capacity (gal)", "Linked To", "Link Type", "Driver", "Ownership", "Status"]
    for ci, h in enumerate(heads, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = hf; c.fill = hfill; c.alignment = center; c.border = border

    row_idx = 2
    for v in vehicles:
        vals = [
            row_idx - 1,
            v["plate_no"],
            v.get("vehicle_category") or "Solo",
            v["vehicle_type"],
            v.get("vehicle_sub_type") or "",
            v.get("model") or "",
            v.get("year") or "",
            v.get("vehicle_length") or "",
            v.get("tank_capacity_gal") or 0,
            v.get("linked_plate_no") or "",
            v.get("link_type") or "",
            v.get("driver_name") or "",
            v["ownership_type"],
            v["status"],
        ]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=row_idx, column=ci, value=val)
            c.border = border
        row_idx += 1

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 8
    ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 16
    ws.column_dimensions["J"].width = 14
    ws.column_dimensions["K"].width = 16
    ws.column_dimensions["L"].width = 24
    ws.column_dimensions["M"].width = 14
    ws.column_dimensions["N"].width = 12

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"vehicles_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(buf, as_attachment=True, download_name=fname, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ── Add Vehicle ─────────────────────────────────────────────────

@fleet_bp.route("/fleet/vehicles/add", methods=["GET", "POST"])
@_login_required("admin")
def vehicle_add():
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    drivers = _all_employees_drivers()
    vehicles_list = db.execute("SELECT plate_no, vehicle_type, model, vehicle_category FROM vehicles WHERE status = 'Active' ORDER BY plate_no").fetchall()
    already_linked = [r["plate_no"] for r in db.execute("SELECT linked_plate_no AS plate_no FROM vehicles WHERE linked_plate_no IS NOT NULL AND linked_plate_no != ''").fetchall()]

    if request.method == "POST":
        plate_no = request.form.get("plate_no", "").strip().upper()
        vehicle_type = request.form.get("vehicle_type", "").strip()
        if vehicle_type == "__custom__":
            vehicle_type = request.form.get("vehicle_type_custom", "").strip()
        model = request.form.get("model", "").strip()
        year = request.form.get("year", "").strip()
        ownership_type = request.form.get("ownership_type", "").strip()
        if ownership_type == "__custom__":
            ownership_type = request.form.get("ownership_type_custom", "").strip()
        partner_name = request.form.get("partner_name", "").strip()
        partner_percent = request.form.get("partner_percent", "").strip()
        driver_id = request.form.get("driver_id", "").strip()
        notes = request.form.get("notes", "").strip()
        vehicle_category = request.form.get("vehicle_category", "Solo").strip()
        if vehicle_category == "__custom__":
            vehicle_category = request.form.get("vehicle_category_custom", "").strip()
        vehicle_sub_type = request.form.get("vehicle_sub_type", "").strip()
        if vehicle_sub_type == "__custom__":
            vehicle_sub_type = request.form.get("vehicle_sub_type_custom", "").strip()
        vehicle_length = request.form.get("vehicle_length", "").strip()
        if vehicle_length == "__custom__":
            vehicle_length = request.form.get("vehicle_length_custom", "").strip()
        tank_capacity_gal = request.form.get("tank_capacity_gal", "0").strip()
        if tank_capacity_gal == "__custom__":
            tank_capacity_gal = request.form.get("tank_capacity_gal_custom", "0").strip()
        linked_plate_no = request.form.get("linked_plate_no", "").strip()
        link_type = request.form.get("link_type", "").strip()

        if not plate_no or not vehicle_type:
            flash("Plate number and vehicle type are required.", "error")
            return render_template("fleet/vehicle_form.html", v=request.form, drivers=drivers, vehicles_list=vehicles_list, vehicle_types=VEHICLE_TYPES, ownership_types=OWNERSHIP_TYPES, vehicle_categories=VEHICLE_CATEGORIES, vehicle_sub_types=VEHICLE_SUB_TYPES, link_types=LINK_TYPES, tank_capacities=TANK_CAPACITIES, already_linked=already_linked, page_title="Add Vehicle", submit_label="Add Vehicle")

        existing = db.execute("SELECT plate_no FROM vehicles WHERE plate_no = ?", (plate_no,)).fetchone()
        if existing:
            flash(f"Vehicle {plate_no} already exists.", "error")
            return render_template("fleet/vehicle_form.html", v=request.form, drivers=drivers, vehicles_list=vehicles_list, vehicle_types=VEHICLE_TYPES, ownership_types=OWNERSHIP_TYPES, vehicle_categories=VEHICLE_CATEGORIES, vehicle_sub_types=VEHICLE_SUB_TYPES, link_types=LINK_TYPES, tank_capacities=TANK_CAPACITIES, already_linked=already_linked, page_title="Add Vehicle", submit_label="Add Vehicle")

        db.execute(
            "INSERT INTO vehicles (plate_no, vehicle_type, model, year, ownership_type, partner_name, partner_percent, status, notes, vehicle_category, vehicle_sub_type, vehicle_length, tank_capacity_gal, linked_plate_no, link_type) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (plate_no, vehicle_type, model, int(year) if year else None, ownership_type, partner_name if ownership_type == "Partnership" else None, float(partner_percent) if partner_percent and ownership_type == "Partnership" else None, "Active", notes, vehicle_category, vehicle_sub_type, vehicle_length, int(tank_capacity_gal) if tank_capacity_gal else 0, linked_plate_no if linked_plate_no else None, link_type if link_type else ''),
        )
        db.commit()

        if driver_id:
            db.execute(
                "INSERT INTO vehicle_assignments (vehicle_id, driver_id, assigned_from, is_current) VALUES (?,?,?,1)",
                (plate_no, driver_id, date.today().isoformat()),
            )
            db.commit()

        flash(f"Vehicle {plate_no} added.", "success")
        return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no))

    return render_template("fleet/vehicle_form.html", v={}, drivers=drivers, vehicles_list=vehicles_list, vehicle_types=VEHICLE_TYPES, ownership_types=OWNERSHIP_TYPES, vehicle_categories=VEHICLE_CATEGORIES, vehicle_sub_types=VEHICLE_SUB_TYPES, link_types=LINK_TYPES, tank_capacities=TANK_CAPACITIES, already_linked=already_linked, page_title="Add Vehicle", submit_label="Add Vehicle")


# ── Edit Vehicle ────────────────────────────────────────────────

@fleet_bp.route("/fleet/vehicles/<path:plate_no>/edit", methods=["GET", "POST"])
@_login_required("admin")
def vehicle_edit(plate_no):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    try:
        v = db.execute("SELECT plate_no, vehicle_type, model, year, ownership_type, partner_name, partner_percent, status, notes, vehicle_category, vehicle_sub_type, vehicle_length, tank_capacity_gal, linked_plate_no, link_type FROM vehicles WHERE plate_no = ?", (plate_no,)).fetchone()
    except Exception as e:
        flash(f"Database error: {e}", "error")
        return redirect(url_for("fleet.vehicle_list"))
    if not v:
        flash("Vehicle not found.", "error")
        return redirect(url_for("fleet.vehicle_list"))
    try:
        drivers = _all_employees_drivers()
    except Exception as e:
        flash(f"Error loading drivers: {e}", "error")
        return redirect(url_for("fleet.vehicle_list"))
    vehicles_list = db.execute("SELECT plate_no, vehicle_type, model, vehicle_category FROM vehicles WHERE status = 'Active' AND plate_no != ? ORDER BY plate_no", (plate_no,)).fetchall()
    linked_trailers = db.execute("SELECT plate_no, vehicle_type, vehicle_sub_type, vehicle_length, tank_capacity_gal FROM vehicles WHERE linked_plate_no = ? AND vehicle_category = 'Trailer' ORDER BY plate_no", (plate_no,)).fetchall()
    # Get already linked vehicle plate_nos (to disable in dropdown)
    already_linked = [r["plate_no"] for r in db.execute("SELECT linked_plate_no AS plate_no FROM vehicles WHERE linked_plate_no IS NOT NULL AND linked_plate_no != ''").fetchall()]

    existing_mulkiyas = db.execute(
        "SELECT id, doc_name, doc_ref_no, expiry_date FROM documents WHERE entity_type = 'vehicle' AND entity_id = ? AND doc_category = 'Mulkiya' ORDER BY uploaded_at DESC",
        (plate_no,)
    ).fetchall()
    from ..documents.routes import _expiry_status
    for doc in existing_mulkiyas:
        doc["_status"] = _expiry_status(doc["expiry_date"])

    if request.method == "POST":
        new_plate = request.form.get("plate_no", "").strip().upper()
        vehicle_type = request.form.get("vehicle_type", "").strip()
        if vehicle_type == "__custom__":
            vehicle_type = request.form.get("vehicle_type_custom", "").strip()
        model = request.form.get("model", "").strip()
        year = request.form.get("year", "").strip()
        ownership_type = request.form.get("ownership_type", "").strip()
        if ownership_type == "__custom__":
            ownership_type = request.form.get("ownership_type_custom", "").strip()
        partner_name = request.form.get("partner_name", "").strip()
        partner_percent = request.form.get("partner_percent", "").strip()
        status = request.form.get("status", "").strip()
        notes = request.form.get("notes", "").strip()
        vehicle_category = request.form.get("vehicle_category", "Solo").strip()
        if vehicle_category == "__custom__":
            vehicle_category = request.form.get("vehicle_category_custom", "").strip()
        vehicle_sub_type = request.form.get("vehicle_sub_type", "").strip()
        if vehicle_sub_type == "__custom__":
            vehicle_sub_type = request.form.get("vehicle_sub_type_custom", "").strip()
        vehicle_length = request.form.get("vehicle_length", "").strip()
        if vehicle_length == "__custom__":
            vehicle_length = request.form.get("vehicle_length_custom", "").strip()
        tank_capacity_gal = request.form.get("tank_capacity_gal", "0").strip()
        if tank_capacity_gal == "__custom__":
            tank_capacity_gal = request.form.get("tank_capacity_gal_custom", "0").strip()
        linked_plate_no = request.form.get("linked_plate_no", "").strip()
        link_type = request.form.get("link_type", "").strip()

        if not new_plate:
            flash("Plate number is required.", "error")
            return render_template("fleet/vehicle_form.html", v=v, drivers=drivers, vehicles_list=vehicles_list, vehicle_types=VEHICLE_TYPES, ownership_types=OWNERSHIP_TYPES, vehicle_categories=VEHICLE_CATEGORIES, vehicle_sub_types=VEHICLE_SUB_TYPES, link_types=LINK_TYPES, tank_capacities=TANK_CAPACITIES, linked_trailers=linked_trailers, existing_mulkiyas=existing_mulkiyas, already_linked=already_linked, page_title="Edit Vehicle", submit_label="Save Changes")

        if new_plate != plate_no:
            existing = db.execute("SELECT plate_no FROM vehicles WHERE plate_no = ?", (new_plate,)).fetchone()
            if existing:
                flash(f"Plate number {new_plate} already exists.", "error")
                return render_template("fleet/vehicle_form.html", v=v, drivers=drivers, vehicles_list=vehicles_list, vehicle_types=VEHICLE_TYPES, ownership_types=OWNERSHIP_TYPES, vehicle_categories=VEHICLE_CATEGORIES, vehicle_sub_types=VEHICLE_SUB_TYPES, link_types=LINK_TYPES, tank_capacities=TANK_CAPACITIES, linked_trailers=linked_trailers, existing_mulkiyas=existing_mulkiyas, already_linked=already_linked, page_title="Edit Vehicle", submit_label="Save Changes")
            try:
                db.execute(
                    "INSERT INTO vehicles (plate_no, vehicle_type, model, year, ownership_type, partner_name, partner_percent, status, notes, vehicle_category, vehicle_sub_type, vehicle_length, tank_capacity_gal, linked_plate_no, link_type) SELECT ?, vehicle_type, model, year, ownership_type, partner_name, partner_percent, status, notes, vehicle_category, vehicle_sub_type, vehicle_length, tank_capacity_gal, linked_plate_no, link_type FROM vehicles WHERE plate_no=?",
                    (new_plate, plate_no),
                )
                db.execute(
                    "UPDATE maintenance_jobs SET vehicle_id=? WHERE vehicle_id=?",
                    (new_plate, plate_no),
                )
                db.execute(
                    "UPDATE vehicle_assignments SET vehicle_id=? WHERE vehicle_id=?",
                    (new_plate, plate_no),
                )
                db.execute("DELETE FROM vehicles WHERE plate_no=?", (plate_no,))
            except Exception as e:
                flash(f"Could not update plate number: {e}", "error")
                return render_template("fleet/vehicle_form.html", v=v, drivers=drivers, vehicles_list=vehicles_list, vehicle_types=VEHICLE_TYPES, ownership_types=OWNERSHIP_TYPES, vehicle_categories=VEHICLE_CATEGORIES, vehicle_sub_types=VEHICLE_SUB_TYPES, link_types=LINK_TYPES, tank_capacities=TANK_CAPACITIES, linked_trailers=linked_trailers, existing_mulkiyas=existing_mulkiyas, already_linked=already_linked, page_title="Edit Vehicle", submit_label="Save Changes")
        else:
            db.execute(
                "UPDATE vehicles SET vehicle_type=?, model=?, year=?, ownership_type=?, partner_name=?, partner_percent=?, status=?, notes=?, vehicle_category=?, vehicle_sub_type=?, vehicle_length=?, tank_capacity_gal=?, linked_plate_no=?, link_type=? WHERE plate_no=?",
                (vehicle_type, model, int(year) if year else None, ownership_type, partner_name if ownership_type == "Partnership" else None, float(partner_percent) if partner_percent and ownership_type == "Partnership" else None, status, notes, vehicle_category, vehicle_sub_type, vehicle_length, int(tank_capacity_gal) if tank_capacity_gal else 0, linked_plate_no if vehicle_category == "Trailer" else None, link_type if vehicle_category == "Trailer" else '', plate_no),
            )
        db.commit()
        flash("Vehicle updated.", "success")
        return redirect(url_for("fleet.vehicle_profile", plate_no=new_plate))

    return render_template("fleet/vehicle_form.html", v=v, drivers=drivers, vehicles_list=vehicles_list, vehicle_types=VEHICLE_TYPES, ownership_types=OWNERSHIP_TYPES, vehicle_categories=VEHICLE_CATEGORIES, vehicle_sub_types=VEHICLE_SUB_TYPES, link_types=LINK_TYPES, tank_capacities=TANK_CAPACITIES, linked_trailers=linked_trailers, existing_mulkiyas=existing_mulkiyas, already_linked=already_linked, page_title="Edit Vehicle", submit_label="Save Changes")


# ── Add Trailer to Head Vehicle ────────────────────────────────

@fleet_bp.route("/fleet/vehicles/<path:plate_no>/add-trailer", methods=["POST"])
@_login_required("admin")
def vehicle_add_trailer(plate_no):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()

    head = db.execute("SELECT plate_no, vehicle_category FROM vehicles WHERE plate_no = ?", (plate_no,)).fetchone()
    if not head:
        flash("Head vehicle not found.", "error")
        return redirect(url_for("fleet.vehicle_list"))
    if (head.get("vehicle_category") or "Solo") != "Head":
        flash("Only Head vehicles can have linked trailers.", "error")
        return redirect(url_for("fleet.vehicle_edit", plate_no=plate_no))

    trailer_plate = request.form.get("trailer_plate_no", "").strip().upper()
    trailer_type = request.form.get("trailer_type", "").strip()
    trailer_sub_type = request.form.get("trailer_sub_type", "").strip()
    trailer_length = request.form.get("trailer_length", "").strip()
    trailer_capacity = request.form.get("trailer_capacity", "0").strip()
    trailer_link_type = request.form.get("trailer_link_type", "Tractor-Flat").strip()

    if not trailer_plate or not trailer_type:
        flash("Trailer plate and type are required.", "error")
        return redirect(url_for("fleet.vehicle_edit", plate_no=plate_no))

    existing = db.execute("SELECT plate_no FROM vehicles WHERE plate_no = ?", (trailer_plate,)).fetchone()
    if existing:
        flash(f"Vehicle {trailer_plate} already exists.", "error")
        return redirect(url_for("fleet.vehicle_edit", plate_no=plate_no))

    db.execute(
        "INSERT INTO vehicles (plate_no, vehicle_type, vehicle_category, vehicle_sub_type, vehicle_length, tank_capacity_gal, linked_plate_no, link_type, status) VALUES (?,?,?,?,?,?,?,?,'Active')",
        (trailer_plate, trailer_type, "Trailer", trailer_sub_type, trailer_length, int(trailer_capacity) if trailer_capacity else 0, plate_no, trailer_link_type),
    )
    db.commit()

    flash(f"Trailer {trailer_plate} linked to {plate_no}.", "success")
    return redirect(url_for("fleet.vehicle_edit", plate_no=plate_no))


# ── Remove Trailer from Head Vehicle ───────────────────────────

@fleet_bp.route("/fleet/vehicles/<path:plate_no>/remove-trailer/<path:trailer_plate>", methods=["POST"])
@_login_required("admin")
def vehicle_remove_trailer(plate_no, trailer_plate):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()

    db.execute("DELETE FROM vehicles WHERE plate_no = ? AND linked_plate_no = ? AND vehicle_category = 'Trailer'", (trailer_plate, plate_no))
    db.commit()

    flash(f"Trailer {trailer_plate} unlinked from {plate_no}.", "success")
    return redirect(url_for("fleet.vehicle_edit", plate_no=plate_no))


# ── Upload Mulkiya for Vehicle ──────────────────────────────────

@fleet_bp.route("/fleet/vehicles/<path:plate_no>/upload-mulkiya", methods=["POST"])
@_login_required("admin")
def vehicle_upload_mulkiya(plate_no):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()

    vehicle = db.execute("SELECT plate_no FROM vehicles WHERE plate_no = ?", (plate_no,)).fetchone()
    if not vehicle:
        flash("Vehicle not found.", "error")
        return redirect(url_for("fleet.vehicle_list"))

    import base64
    doc_name = request.form.get("doc_name", "").strip()
    doc_ref_no = request.form.get("doc_ref_no", "").strip() or None
    expiry_date = request.form.get("expiry_date", "").strip() or None
    file = request.files.get("file")

    if not doc_name or not file:
        flash("Document name and file are required.", "error")
        return redirect(url_for("fleet.vehicle_edit", plate_no=plate_no))

    file_data = base64.b64encode(file.read()).decode("utf-8")
    file_type = file.content_type or "application/octet-stream"
    file_size = len(file_data)

    from ..documents.routes import _generate_thumbnail
    thumbnail_data, pdf_preview_data = _generate_thumbnail(file_data, file_type)

    db.execute(
        """INSERT INTO documents (entity_type, entity_id, doc_name, doc_category, doc_ref_no,
           issue_date, expiry_date, file_data, file_type, file_size, thumbnail_data, pdf_preview_data)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
        ("vehicle", plate_no, doc_name, "Mulkiya", doc_ref_no,
         None, expiry_date, file_data, file_type, file_size, thumbnail_data, pdf_preview_data),
    )
    db.commit()
    db.close()

    flash(f"Mulkiya '{doc_name}' uploaded for {plate_no}.", "success")
    return redirect(url_for("fleet.vehicle_edit", plate_no=plate_no))


# ── Remove Mulkiya from Vehicle ─────────────────────────────────

@fleet_bp.route("/fleet/vehicles/<path:plate_no>/remove-mulkiya/<int:doc_id>", methods=["POST"])
@_login_required("admin")
def vehicle_remove_mulkiya(plate_no, doc_id):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()

    db.execute("DELETE FROM documents WHERE id = ? AND entity_type = 'vehicle' AND entity_id = ?", (doc_id, plate_no))
    db.commit()
    db.close()

    flash("Mulkiya deleted.", "success")
    return redirect(url_for("fleet.vehicle_edit", plate_no=plate_no))


# ── Unlink Vehicle from Head ────────────────────────────────────

@fleet_bp.route("/fleet/vehicles/<path:plate_no>/unlink", methods=["POST"])
@_login_required("admin")
def vehicle_unlink(plate_no):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()

    v = db.execute("SELECT plate_no, vehicle_category FROM vehicles WHERE plate_no = ?", (plate_no,)).fetchone()
    if not v:
        flash("Vehicle not found.", "error")
        return redirect(url_for("fleet.vehicle_list"))

    db.execute("UPDATE vehicles SET linked_plate_no = NULL, link_type = '' WHERE plate_no = ?", (plate_no,))
    db.commit()
    db.close()

    flash(f"Vehicle {plate_no} unlinked from Head.", "success")
    return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no))


# ── Link Vehicle to Head ────────────────────────────────────────

@fleet_bp.route("/fleet/vehicles/<path:plate_no>/link-to-head", methods=["POST"])
@_login_required("admin")
def vehicle_link_to_head(plate_no):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()

    v = db.execute("SELECT plate_no, vehicle_category FROM vehicles WHERE plate_no = ?", (plate_no,)).fetchone()
    if not v:
        flash("Vehicle not found.", "error")
        return redirect(url_for("fleet.vehicle_list"))

    head_plate = request.form.get("head_plate_no", "").strip()
    link_type = request.form.get("link_type", "").strip()

    if not head_plate:
        flash("Please select a Head vehicle.", "error")
        return redirect(url_for("fleet.vehicle_edit", plate_no=plate_no))

    head = db.execute("SELECT plate_no, vehicle_category FROM vehicles WHERE plate_no = ?", (head_plate,)).fetchone()
    if not head:
        flash("Head vehicle not found.", "error")
        return redirect(url_for("fleet.vehicle_edit", plate_no=plate_no))

    db.execute("UPDATE vehicles SET linked_plate_no = ?, link_type = ? WHERE plate_no = ?", (head_plate, link_type, plate_no))
    db.commit()
    db.close()

    flash(f"Vehicle {plate_no} linked to Head {head_plate}.", "success")
    return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no))


# ── Vehicle Profile ─────────────────────────────────────────────

@fleet_bp.route("/fleet/vehicles/<path:plate_no>")
@_login_required("admin")
def vehicle_profile(plate_no):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    v = _vehicle_full(plate_no)
    if not v:
        flash("Vehicle not found.", "error")
        return redirect(url_for("fleet.vehicle_list"))

    # Linked vehicles (vehicles linked TO this Head)
    linked_vehicles = db.execute(
        "SELECT plate_no, vehicle_type, model, linked_plate_no, link_type FROM vehicles WHERE linked_plate_no = ? ORDER BY plate_no",
        (plate_no,),
    ).fetchall()

    # Current drivers (latest assignment per vehicle, is_current=1)
    current_drivers = db.execute(
        """SELECT DISTINCT ON (va.vehicle_id) va.vehicle_id, va.driver_id, va.assigned_from, e.full_name AS driver_name
           FROM vehicle_assignments va
           JOIN employees e ON e.employee_id = va.driver_id
           WHERE va.vehicle_id = ? AND va.is_current = 1
           ORDER BY va.vehicle_id, va.assigned_from DESC""",
        (plate_no,),
    ).fetchall()

    active_tab = request.args.get("tab", "overview")
    highlight = request.args.get("highlight", "")

    # Driver history
    driver_history = db.execute(
        """SELECT va.*, e.full_name AS driver_name FROM vehicle_assignments va
           JOIN employees e ON e.employee_id = va.driver_id
           WHERE va.vehicle_id = ? ORDER BY va.assigned_from DESC""",
        (plate_no,),
    ).fetchall()

    # Approved jobs (maintenance_jobs + maintenance_papers)
    approved_jobs = db.execute(
        f"""SELECT {_MJ_LIST_COLS}, mj.tax_amount AS vat_amount, (mj.amount - mj.tax_amount) AS net_amount, COALESCE(s.full_name, 'Admin') AS staff_name FROM maintenance_jobs mj
           LEFT JOIN field_staff s ON s.staff_id = mj.staff_id
           WHERE mj.vehicle_id = ? AND mj.status = 'approved'
           ORDER BY mj.created_at DESC""",
        (plate_no,),
    ).fetchall()

    raw_papers = db.execute(
        """SELECT mp.paper_no AS id, vm.vehicle_no AS vehicle_id, mp.vehicle_id AS edit_vehicle_id,
                  mp.technician_code AS staff_id,
                  mp.total_amount AS amount, mp.work_summary AS description,
                  mp.review_status AS status, mp.notes AS admin_notes,
                  mp.attachment_path AS attachment_name, mp.created_at,
                  'Maintenance' AS category, '' AS attachment_type,
                  NULL AS attachment_data,
                  mp.tax_mode AS tax_mode, mp.supplier_name AS supplier_name,
                  mp.supplier_trn AS supplier_trn, mp.subtotal AS subtotal,
                  mp.tax_amount AS vat_amount, mp.total_amount AS total_amount,
                  COALESCE(s.full_name, '') AS staff_name
           FROM maintenance_papers mp
           JOIN vehicle_master vm ON vm.vehicle_id = mp.vehicle_id
           LEFT JOIN field_staff s ON s.staff_id = mp.technician_code
           WHERE vm.vehicle_no = ?
             AND mp.review_status IN ('Approved', 'Pending')
           ORDER BY mp.created_at DESC""",
        (plate_no,),
    ).fetchall()
    maintenance_papers_list = []
    for r in raw_papers:
        d = dict(r)
        if isinstance(d.get("created_at"), datetime):
            d["created_at"] = d["created_at"].strftime("%Y-%m-%d %H:%M:%S")
        maintenance_papers_list.append(d)

    approved_jobs_list = []
    for r in approved_jobs:
        d = dict(r)
        if isinstance(d.get("created_at"), datetime):
            d["created_at"] = d["created_at"].strftime("%Y-%m-%d %H:%M:%S")
        approved_jobs_list.append(d)

    combined = sorted(
        approved_jobs_list + maintenance_papers_list,
        key=lambda x: (x.get("created_at") or ""),
        reverse=True,
    )

    # Fetch ALL vehicle documents from unified documents table
    all_docs = db.execute(
        "SELECT id, doc_name, doc_category, file_data AS doc_data, file_type AS doc_type, entity_id, expiry_date, uploaded_at, thumbnail_data, pdf_preview_data, notes FROM documents WHERE entity_type = 'vehicle' AND entity_id = ? ORDER BY uploaded_at DESC",
        (plate_no,),
    ).fetchall()
    from ..documents.routes import _expiry_status
    for md in all_docs:
        md["_status"] = _expiry_status(md.get("expiry_date"))

    fuel_entries = db.execute(
        "SELECT id, vehicle_plate, entry_date, gallons, rate_per_gallon, total_amount, supplier_id, supplier_name, notes, source_expense_id, created_at FROM fuel_entries WHERE vehicle_plate = ? ORDER BY entry_date DESC, id DESC",
        (plate_no,),
    ).fetchall()
    fuel_total_gallons = sum(f["gallons"] for f in fuel_entries) if fuel_entries else 0
    fuel_total_amount = sum(f["total_amount"] for f in fuel_entries) if fuel_entries else 0
    suppliers = db.execute("SELECT id, supplier_name FROM suppliers WHERE status = 'Active' ORDER BY supplier_name").fetchall()

    # Supplier bills (Parts) for this vehicle
    try:
        supplier_bills = db.execute(
            """SELECT sb.*, s.supplier_name FROM supplier_bills sb
               JOIN suppliers s ON s.id = sb.supplier_id
               WHERE sb.vehicle_plate = ? ORDER BY sb.bill_date DESC, sb.id DESC""",
            (plate_no,),
        ).fetchall()
    except Exception:
        supplier_bills = []
    parts_total_amount = sum(b["total_amount"] for b in supplier_bills) if supplier_bills else 0
    parts_total_vat = sum(b["vat_amount"] for b in supplier_bills) if supplier_bills else 0
    parts_total_net = sum(b["net_amount"] for b in supplier_bills) if supplier_bills else 0

    # Traffic fines
    try:
        traffic_fines = db.execute(
            "SELECT * FROM traffic_fines WHERE plate_no = ? ORDER BY fine_date DESC",
            (plate_no,),
        ).fetchall()
    except Exception:
        traffic_fines = []
    fines_total = sum(f["fine_amount"] for f in traffic_fines) if traffic_fines else 0
    fines_paid = sum(f["fine_amount"] for f in traffic_fines if f["fine_status"] in ("Paid", "Closed")) if traffic_fines else 0
    fines_pending = fines_total - fines_paid

    head_vehicles = db.execute(
        "SELECT plate_no, vehicle_type, model FROM vehicles WHERE vehicle_category = 'Head' AND status = 'Active' ORDER BY plate_no"
    ).fetchall()
    already_linked = [r["plate_no"] for r in db.execute("SELECT linked_plate_no AS plate_no FROM vehicles WHERE linked_plate_no IS NOT NULL AND linked_plate_no != ''").fetchall()]

    # Ghadeer card (TAQA prepaid water-filling card) linked to this vehicle
    try:
        ghadeer_card = db.execute(
            """SELECT c.id, c.card_no, c.cardholder, c.status,
                      COALESCE((SELECT SUM(amount) FROM ghadeer_transactions WHERE card_id = c.id), 0) AS balance
               FROM ghadeer_cards c
               WHERE c.vehicle_plate = ? AND c.status = 'Active'
               ORDER BY c.id LIMIT 1""",
            (plate_no,),
        ).fetchone()
    except Exception:
        ghadeer_card = None

    return render_template(
        "fleet/vehicle_profile.html",
        v=v,
        active_tab=active_tab,
        highlight=highlight,
        driver_history=driver_history,
        approved_jobs=approved_jobs_list,
        maintenance_papers_list=maintenance_papers_list,
        combined_jobs=combined,
        all_drivers=_all_employees_drivers(),
        all_docs=all_docs,
        fuel_entries=fuel_entries,
        fuel_total_gallons=fuel_total_gallons,
        fuel_total_amount=fuel_total_amount,
        supplier_bills=supplier_bills,
        parts_total_amount=parts_total_amount,
        parts_total_vat=parts_total_vat,
        parts_total_net=parts_total_net,
        suppliers=suppliers,
        linked_vehicles=linked_vehicles,
        current_drivers=current_drivers,
        head_vehicles=head_vehicles,
        already_linked=already_linked,
        link_types=LINK_TYPES,
        traffic_fines=traffic_fines,
        fines_total=fines_total,
        fines_paid=fines_paid,
        fines_pending=fines_pending,
        ghadeer_card=ghadeer_card,
        date=date,
    )


# ── Delete Vehicle ────────────────────────────────────────────

@fleet_bp.route("/fleet/vehicles/<path:plate_no>/delete", methods=["POST"])
@_login_required("admin")
def vehicle_delete(plate_no):
    _touch_admin_workspace("fleet")
    db = open_db()
    v = _vehicle_full(plate_no)
    if not v:
        flash("Vehicle not found.", "error")
        return redirect(url_for("fleet.vehicle_list"))
    db.execute("DELETE FROM vehicle_assignments WHERE vehicle_id = ?", (plate_no,))
    db.execute("DELETE FROM vehicles WHERE plate_no = ?", (plate_no,))
    db.commit()
    flash(f"Vehicle {plate_no} deleted.", "success")
    return redirect(url_for("fleet.vehicle_list"))


# ── Assign/Replace Driver ───────────────────────────────────────

@fleet_bp.route("/fleet/vehicles/<path:plate_no>/assign", methods=["POST"])
@_login_required("admin")
def vehicle_assign_driver(plate_no):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()

    driver_id = request.form.get("driver_id", "").strip()
    assigned_from = request.form.get("assigned_from", "").strip() or date.today().isoformat()

    if not driver_id:
        flash("Please select a driver.", "error")
        return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no))

    # Remove driver from any OTHER vehicle (not this one) — allows day/night shift on same vehicle
    db.execute(
        "UPDATE vehicle_assignments SET is_current = 0, assigned_until = ? WHERE driver_id = ? AND is_current = 1 AND vehicle_id != ?",
        (assigned_from, driver_id, plate_no),
    )
    # Check if this driver is already assigned to this vehicle
    existing = db.execute(
        "SELECT id FROM vehicle_assignments WHERE vehicle_id = ? AND driver_id = ? AND is_current = 1",
        (plate_no, driver_id),
    ).fetchone()
    if not existing:
        # Insert new assignment only if not already assigned
        db.execute(
            "INSERT INTO vehicle_assignments (vehicle_id, driver_id, assigned_from, is_current) VALUES (?,?,?,1)",
            (plate_no, driver_id, assigned_from),
        )
    db.commit()

    flash(f"Driver assigned to {plate_no}.", "success")
    return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no, tab="driver"))


# ── Field Staff: Staff Login ────────────────────────────────────




# ── Vehicle Documents ─────────────────────────────────────────────

@fleet_bp.route("/fleet/vehicles/<path:plate_no>/documents/upload", methods=["POST"])
@_login_required("admin")
def vehicle_document_upload(plate_no):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    v = _vehicle_full(plate_no)
    if not v:
        flash("Vehicle not found.", "error")
        return redirect(url_for("fleet.vehicle_list"))
    doc_name = request.form.get("doc_name", "").strip()
    notes = request.form.get("notes", "").strip()
    if not doc_name:
        flash("Document name is required.", "error")
        return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no))
    import base64
    doc_data = None
    doc_type = None
    file_data = None
    if "doc_file" in request.files:
        f = request.files["doc_file"]
        if f.filename:
            raw = f.read()
            doc_data = base64.b64encode(raw).decode("utf-8")
            doc_type = f.content_type
            file_data = doc_data

    is_mulkiya = "mulkiya" in doc_name.lower()
    doc_category = "Mulkiya" if is_mulkiya else "Other"

    thumbnail_data = None
    pdf_preview_data = None
    if file_data:
        try:
            from ..documents.routes import _generate_thumbnail
            thumbnail_data, pdf_preview_data = _generate_thumbnail(file_data, doc_type)
        except Exception:
            pass
    db.execute(
        "INSERT INTO documents (entity_type, entity_id, doc_name, doc_category, file_data, file_type, file_size, thumbnail_data, pdf_preview_data, notes) VALUES ('vehicle',?,?,?,?,?,?,?,?,?)",
        (plate_no, doc_name, doc_category, file_data, doc_type, len(file_data) if file_data else 0, thumbnail_data, pdf_preview_data, notes),
    )
    db.commit()
    flash(f"Document '{doc_name}' uploaded.", "success")
    return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no, tab='documents'))


@fleet_bp.route("/fleet/vehicles/<path:plate_no>/documents/<int:doc_id>/delete", methods=["POST"])
@_login_required("admin")
def vehicle_document_delete(plate_no, doc_id):
    _touch_admin_workspace("fleet")
    db = open_db()
    db.execute("DELETE FROM documents WHERE id = ? AND entity_type = 'vehicle' AND entity_id = ?", (doc_id, plate_no))
    db.commit()
    flash("Document deleted.", "success")
    return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no, tab='documents'))


@fleet_bp.route("/fleet/vehicles/<path:plate_no>/documents/<int:doc_id>/view")
@_login_required("admin")
def vehicle_document_view(plate_no, doc_id):
    db = open_db()
    doc = db.execute("SELECT id, doc_name, file_type AS doc_type, file_data AS doc_data FROM documents WHERE id = ? AND entity_type = 'vehicle' AND entity_id = ?", (doc_id, plate_no)).fetchone()
    if not doc or not doc["doc_data"]:
        flash("Document not found.", "error")
        return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no))
    import base64
    from io import BytesIO
    data = base64.b64decode(doc["doc_data"])
    return send_file(
        BytesIO(data),
        mimetype=doc["doc_type"] or "application/octet-stream",
        as_attachment=False,
        download_name=doc["doc_name"] or f"document_{doc_id}",
    )


@fleet_bp.route("/fleet/vehicles/<path:plate_no>/documents/<int:doc_id>/edit", methods=["POST"])
@_login_required("admin")
def vehicle_document_edit(plate_no, doc_id):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    v = _vehicle_full(plate_no)
    if not v:
        flash("Vehicle not found.", "error")
        return redirect(url_for("fleet.vehicle_list"))

    doc_name = request.form.get("doc_name", "").strip()
    notes = request.form.get("notes", "").strip()
    expiry_date = request.form.get("expiry_date", "").strip() or None

    if not doc_name:
        flash("Document name is required.", "error")
        return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no, tab="documents"))

    db.execute(
        "UPDATE documents SET doc_name = ?, notes = ?, expiry_date = ? WHERE id = ? AND entity_type = 'vehicle' AND entity_id = ?",
        (doc_name, notes, expiry_date, doc_id, plate_no),
    )
    db.commit()
    db.close()
    flash(f"Document '{doc_name}' updated.", "success")
    return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no, tab="documents"))


# ── Vehicle Traffic Fines ────────────────────────────────────────

@fleet_bp.route("/fleet/vehicles/<path:plate_no>/traffic-fines/upload", methods=["POST"])
@_login_required("admin")
def vehicle_traffic_fine_upload(plate_no):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    v = _vehicle_full(plate_no)
    if not v:
        flash("Vehicle not found.", "error")
        return redirect(url_for("fleet.vehicle_list"))

    fine_date = request.form.get("fine_date", "").strip()
    fine_type = request.form.get("fine_type", "Traffic Fine").strip()
    fine_amount = request.form.get("fine_amount", "0").strip()
    fine_location = request.form.get("fine_location", "").strip()
    fine_description = request.form.get("fine_description", "").strip()
    fine_status = request.form.get("fine_status", "Pending").strip()
    notes = request.form.get("notes", "").strip()

    if not fine_date:
        flash("Fine date is required.", "error")
        return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no, tab="fines"))

    try:
        fine_amount = float(fine_amount)
    except ValueError:
        fine_amount = 0

    attachment_data = None
    attachment_type = None
    attachment_name = None
    if "fine_attachment" in request.files:
        f = request.files["fine_attachment"]
        if f.filename:
            import base64
            raw = f.read()
            attachment_data = base64.b64encode(raw).decode("utf-8")
            attachment_type = f.content_type
            attachment_name = f.filename

    db.execute(
        """INSERT INTO traffic_fines (plate_no, fine_date, fine_type, fine_amount, fine_location, fine_description, fine_status, attachment_data, attachment_type, attachment_name, notes)
           VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
        (plate_no, fine_date, fine_type, fine_amount, fine_location, fine_description, fine_status, attachment_data, attachment_type, attachment_name, notes),
    )
    db.commit()
    flash(f"Traffic fine of AED {fine_amount:.2f} recorded.", "success")
    return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no, tab="fines"))


@fleet_bp.route("/fleet/vehicles/<path:plate_no>/traffic-fines/<int:fine_id>/delete", methods=["POST"])
@_login_required("admin")
def vehicle_traffic_fine_delete(plate_no, fine_id):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    db.execute("DELETE FROM traffic_fines WHERE id = ? AND plate_no = ?", (fine_id, plate_no))
    db.commit()
    flash("Traffic fine deleted.", "success")
    return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no, tab="fines"))


@fleet_bp.route("/fleet/vehicles/<path:plate_no>/traffic-fines/<int:fine_id>/attachment")
@_login_required("admin")
def vehicle_traffic_fine_attachment(plate_no, fine_id):
    ensure_fleet_tables()
    db = open_db()
    fine = db.execute("SELECT attachment_data, attachment_type, attachment_name FROM traffic_fines WHERE id = ? AND plate_no = ?", (fine_id, plate_no)).fetchone()
    db.close()
    if not fine or not fine["attachment_data"]:
        flash("Attachment not found.", "error")
        return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no, tab="fines"))
    import base64
    data = base64.b64decode(fine["attachment_data"])
    return Response(data, mimetype=fine["attachment_type"] or "application/octet-stream",
                    headers={"Content-Disposition": f'inline; filename="{fine["attachment_name"] or "attachment"}"'})



# ═════════════════════════════════════════════════════════════════
# ADMIN: Field Staff Management
# ═════════════════════════════════════════════════════════════════

def _staff_photo_url(row):
    if row and row.get("photo_data") and row.get("photo_content_type"):
        return f"data:{row['photo_content_type']};base64,{row['photo_data']}"
    return None


def _sync_field_staff_to_technician(db, staff_id, full_name, phone, username, pw_hash, is_active):
    status = "Active" if is_active else "Inactive"

    existing = db.execute(
        "SELECT technician_code FROM technicians WHERE technician_code = ?",
        (staff_id,),
    ).fetchone()
    if existing:
        db.execute("""
            UPDATE technicians
            SET user_id = ?, password_hash = ?, phone_number = ?,
                specialization = ?, status = ?
            WHERE technician_code = ?
        """, (username, pw_hash, phone, full_name, status, staff_id))
        return

    user_taken = db.execute(
        "SELECT technician_code FROM technicians WHERE user_id = ?",
        (username,),
    ).fetchone()
    if user_taken:
        db.execute("""
            UPDATE technicians
            SET technician_code = ?, password_hash = ?, phone_number = ?,
                specialization = ?, status = ?
            WHERE user_id = ?
        """, (staff_id, pw_hash, phone, full_name, status, username))
        return

    db.execute("""
        INSERT INTO technicians
        (technician_code, party_code, user_id, password_hash, phone_number, specialization, status)
        VALUES (?, NULL, ?, ?, ?, ?, ?)
    """, (staff_id, username, pw_hash, phone, full_name, status))


def _import_field_staff_from_sqlite(db):
    backend = current_app.config.get("DATABASE_BACKEND", "sqlite")
    if backend != "postgres":
        return
    existing = db.execute("SELECT COUNT(*) AS c FROM field_staff").fetchone()["c"] or 0
    if existing > 0:
        return
    try:
        sqlite_path = Path(current_app.config.get("DATABASE", "payroll.db"))
        if not sqlite_path.exists():
            sqlite_path = Path(current_app.root_path).parent / "payroll.db"
        if not sqlite_path.exists():
            return
        sdb = sqlite3.connect(str(sqlite_path))
        sdb.row_factory = sqlite3.Row
    except Exception:
        return

    try:
        old_staff = sdb.execute("SELECT staff_id, full_name, phone, username, password_hash, is_active FROM field_staff").fetchall()
    except Exception:
        old_staff = []

    for s in old_staff:
        try:
            pw_hash = s["password_hash"] or generate_password_hash("changeme123")
            db.execute(
                """INSERT INTO field_staff (staff_id, full_name, phone, username, password_hash, is_active, created_at)
                   VALUES (?, ?, ?, ?, ?, ?, COALESCE(?, CURRENT_TIMESTAMP))""",
                (s["staff_id"], s["full_name"], s["phone"] or "", s["username"],
                 pw_hash, s["is_active"], s.get("created_at")),
            )
        except Exception:
            pass
    if old_staff:
        db.commit()

    try:
        old_jobs = sdb.execute("SELECT id, vehicle_id, staff_id, amount, category, description, attachment_name, attachment_data, attachment_type, status, admin_notes, approved_at, supplier_name, supplier_trn, supplier_bill_no, tax_mode, tax_amount FROM maintenance_jobs").fetchall()
    except Exception:
        try:
            old_jobs = sdb.execute("SELECT id, vehicle_id, staff_id, amount, category, description, attachment_name, attachment_data, attachment_type, status, admin_notes, approved_at, supplier_name, supplier_trn, tax_mode, tax_amount FROM maintenance_jobs").fetchall()
        except Exception:
            try:
                old_jobs = sdb.execute("SELECT id, vehicle_id, staff_id, amount, category, description, attachment_name, attachment_data, attachment_type, status, admin_notes, approved_at FROM maintenance_jobs").fetchall()
            except Exception:
                old_jobs = []

    existing_papers = set()
    try:
        rows = db.execute("SELECT paper_no FROM maintenance_papers").fetchall()
        existing_papers = {r["paper_no"] for r in rows}
    except Exception:
        pass

    for j in old_jobs:
        pno = f"PAPER-{j['id']:04d}"
        if pno in existing_papers:
            continue
        status_map = {"pending": "Pending", "approved": "Approved", "rejected": "Rejected"}
        rev_status = status_map.get(j["status"], "Pending")
        paper_date = (j["created_at"] or "")[:10] or "2025-01-01"
        try:
            tax_mode = j.get("tax_mode") or "Without Tax"
            amount = float(j["amount"] or 0)
            tax_amount = float(j.get("tax_amount") or 0) if tax_mode == "Tax Invoice" else 0.0
            subtotal = amount - tax_amount
            db.execute("""
                INSERT INTO maintenance_papers
                (paper_no, paper_date, vehicle_id, technician_code, work_summary,
                 total_amount, tax_mode, subtotal, tax_amount, supplier_name, supplier_trn, supplier_bill_no,
                 review_status, payment_status, notes, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Pending', ?, ?)
            """, (pno, paper_date, j["vehicle_id"], j["staff_id"],
                  j["description"] or "", amount, tax_mode, subtotal, tax_amount,
                  j.get("supplier_name") or "", j.get("supplier_trn") or "", j.get("supplier_bill_no") or "",
                  rev_status, j["admin_notes"] or "", j["created_at"]))
            existing_papers.add(pno)
        except Exception:
            pass
    if old_jobs:
        db.commit()

    try:
        sdb.close()
    except Exception:
        pass


def _import_maintenance_staff_from_sqlite(db):
    backend = current_app.config.get("DATABASE_BACKEND", "sqlite")
    if backend != "postgres":
        return
    try:
        sqlite_path = Path(current_app.config.get("DATABASE", "payroll.db"))
        if not sqlite_path.exists():
            sqlite_path = Path(current_app.root_path).parent / "payroll.db"
        if not sqlite_path.exists():
            return
        sdb = sqlite3.connect(str(sqlite_path))
        sdb.row_factory = sqlite3.Row
    except Exception:
        return
    try:
        old_staff = sdb.execute("SELECT id, staff_id, full_name, phone, role, status, created_at FROM maintenance_staff").fetchall()
    except Exception:
        old_staff = []
    for s in old_staff:
        code = s["staff_id"]
        existing = db.execute("SELECT staff_code FROM maintenance_staff WHERE staff_code = ?", (code,)).fetchone()
        if existing:
            continue
        try:
            db.execute("""
                INSERT INTO maintenance_staff (staff_code, staff_name, phone_number, status, notes, created_at)
                VALUES (?, ?, ?, ?, ?, COALESCE(?, CURRENT_TIMESTAMP))
            """, (code, s["full_name"], s["phone"] or "", s["status"] or "Active",
                  "", s["created_at"]))
        except Exception:
            pass
        already = db.execute("SELECT staff_id FROM field_staff WHERE staff_id = ?", (code,)).fetchone()
        if already:
            continue
        name = s["full_name"]
        username = (code + name)[:20].lower().replace("-", "").replace(" ", "")
        pw_hash = generate_password_hash("changeme123")
        try:
            db.execute("""
                INSERT INTO field_staff (staff_id, full_name, phone, username, password_hash, is_active)
                VALUES (?, ?, ?, ?, ?, 1)
            """, (code, name, s["phone"] or "", username, pw_hash))
        except Exception:
            continue
        try:
            db.execute("""
                INSERT INTO technicians (technician_code, party_code, user_id, password_hash, phone_number, specialization, status)
                VALUES (?, NULL, ?, ?, ?, ?, 'Active')
            """, (code, username, pw_hash, s["phone"] or "", name))
        except Exception:
            pass
    db.execute("""
        UPDATE maintenance_papers mp
        SET technician_code = fs.staff_id
        FROM field_staff fs
        WHERE mp.technician_code IS NULL
        AND mp.staff_code IS NOT NULL
        AND mp.staff_code = fs.staff_id
    """)
    if old_staff:
        db.commit()
    try:
        sdb.close()
    except Exception:
        pass


def _import_orphaned_maintenance_jobs(db):
    orphan_staff = db.execute("""
        SELECT DISTINCT mj.staff_id FROM maintenance_jobs mj
        LEFT JOIN field_staff fs ON fs.staff_id = mj.staff_id
        WHERE fs.staff_id IS NULL
    """).fetchall()
    for row in orphan_staff:
        staff_id = row["staff_id"]
        sample = db.execute(
            "SELECT mj.id FROM maintenance_jobs mj WHERE mj.staff_id = ? LIMIT 1",
            (staff_id,),
        ).fetchone()
        if not sample:
            continue
        name = f"Staff {staff_id}"
        username = staff_id.lower()
        pw_hash = generate_password_hash("changeme123")
        try:
            db.execute("""
                INSERT INTO field_staff (staff_id, full_name, phone, username, password_hash, is_active)
                VALUES (?, ?, '', ?, ?, 1)
            """, (staff_id, name, username, pw_hash))
        except Exception:
            continue
        try:
            db.execute("""
                INSERT INTO technicians (technician_code, party_code, user_id, password_hash, phone_number, specialization, status)
                VALUES (?, NULL, ?, ?, '', ?, 'Active')
            """, (staff_id, username, pw_hash, name))
        except Exception:
            pass
    if orphan_staff:
        db.commit()


@fleet_bp.route("/fleet/maintenance-entry", methods=["GET", "POST"])
@_login_required("admin")
def fleet_maintenance_entry():
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    vehicles = db.execute("SELECT plate_no, vehicle_type, model FROM vehicles ORDER BY plate_no").fetchall()
    if not db.execute("SELECT staff_id FROM field_staff WHERE staff_id='admin'").fetchone():
        try:
            db.execute("INSERT INTO field_staff (staff_id, full_name, username, password_hash, phone, is_active) VALUES ('admin','System Admin','admin','',NULL,1)")
            db.commit()
        except Exception:
            pass
    if request.method == "POST":
        import base64
        vehicle_ids = request.form.getlist("vehicle_id[]")
        amounts = request.form.getlist("amount[]")
        categories = request.form.getlist("category[]")
        descriptions = request.form.getlist("description[]")
        entry_dates = request.form.getlist("entry_date[]")
        tax_modes = request.form.getlist("tax_mode[]")
        supplier_names = request.form.getlist("supplier_name[]")
        supplier_bill_nos = request.form.getlist("supplier_bill_no[]")
        if not vehicle_ids or not vehicle_ids[0].strip():
            flash("Please select a vehicle.", "error")
            return render_template("fleet/fleet_maintenance_entry.html", vehicles=vehicles, today=date.today().isoformat())
        errors = []
        for i, vid in enumerate(vehicle_ids):
            vid = vid.strip()
            if not vid:
                continue
            tm = tax_modes[i].strip() if i < len(tax_modes) else "Without Tax"
            if tm == "Tax Invoice":
                sn = supplier_names[i].strip() if i < len(supplier_names) else ""
                sbn = supplier_bill_nos[i].strip() if i < len(supplier_bill_nos) else ""
                if not sn:
                    errors.append(f"Entry {i+1} ({vid}): Supplier name is required for Tax Invoice bill.")
                if not sbn:
                    errors.append(f"Entry {i+1} ({vid}): Bill number is required for Tax Invoice bill.")
        if errors:
            flash(" ".join(errors), "error")
            return render_template("fleet/fleet_maintenance_entry.html", vehicles=vehicles, today=date.today().isoformat())
        inserted = 0
        last_vehicle = vehicle_ids[0].strip()
        for i, vid in enumerate(vehicle_ids):
            vid = vid.strip()
            if not vid:
                continue
            last_vehicle = vid
            try:
                amt = float(amounts[i]) if i < len(amounts) and amounts[i].strip() else 0
            except (ValueError, IndexError):
                amt = 0
            cat = categories[i].strip() if i < len(categories) else ""
            desc = descriptions[i].strip() if i < len(descriptions) else ""
            edate = entry_dates[i].strip() if i < len(entry_dates) else ""
            if not edate:
                edate = date.today().isoformat()
            tm = tax_modes[i].strip() if i < len(tax_modes) else "Without Tax"
            sn = supplier_names[i].strip() if i < len(supplier_names) else ""
            sbn = supplier_bill_nos[i].strip() if i < len(supplier_bill_nos) else ""
            tax_amount = 0.0
            staff_amount = amt
            if tm == "Tax Invoice":
                tax_amount = round(amt * 0.05, 2)
                staff_amount = amt
                amt = round(amt + tax_amount, 2)
            att_name = None
            att_data = None
            att_type = None
            att_key = f"attachment_{i}"
            if request.files and att_key in request.files:
                f = request.files[att_key]
                if f and f.filename:
                    att_name = f.filename
                    att_data = base64.b64encode(f.read()).decode("utf-8")
                    att_type = f.content_type or "application/octet-stream"
            db.execute(
                "INSERT INTO maintenance_jobs (vehicle_id, staff_id, amount, category, description, status, created_at, attachment_name, attachment_data, attachment_type, tax_mode, tax_amount, staff_amount, supplier_name, supplier_bill_no) VALUES (?, ?, ?, ?, ?, 'approved', ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (vid, 'admin', amt, cat, desc, edate, att_name, att_data, att_type, tm, tax_amount, staff_amount, sn, sbn)
            )
            inserted += 1
        db.commit()
        db.close()
        if inserted == 1:
            flash(f"Maintenance entry added and approved for vehicle {last_vehicle}.", "success")
        else:
            flash(f"{inserted} maintenance entries added and approved.", "success")
        return redirect(url_for("fleet.fleet_maintenance_entry"))
    return render_template("fleet/fleet_maintenance_entry.html", vehicles=vehicles, today=date.today().isoformat())


@fleet_bp.route("/fleet/vehicles/<path:plate_no>/add-maintenance", methods=["POST"])
@_login_required("admin")
def vehicle_add_maintenance(plate_no):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    v = _vehicle_full(plate_no)
    if not v:
        flash("Vehicle not found.", "error")
        return redirect(url_for("fleet.vehicle_list"))
    import base64
    amount = request.form.get("amount", "0").strip()
    category = request.form.get("category", "").strip()
    description = request.form.get("description", "").strip()
    entry_date = request.form.get("entry_date", "").strip() or date.today().isoformat()
    try:
        amount = float(amount) if amount else 0
    except ValueError:
        amount = 0
    attachment_name = None
    attachment_data = None
    attachment_type = None
    if request.files and "attachment" in request.files:
        f = request.files["attachment"]
        if f and f.filename:
            attachment_name = f.filename
            attachment_data = base64.b64encode(f.read()).decode("utf-8")
            attachment_type = f.content_type or "application/octet-stream"
    db.execute(
        "INSERT INTO maintenance_jobs (vehicle_id, staff_id, amount, category, description, status, created_at, attachment_name, attachment_data, attachment_type) VALUES (?, ?, ?, ?, ?, 'approved', ?, ?, ?, ?)",
        (plate_no, 'admin', amount, category, description, entry_date, attachment_name, attachment_data, attachment_type)
    )
    db.commit()
    db.close()
    flash(f"Maintenance entry added for {plate_no}.", "success")
    return redirect(url_for("fleet.vehicle_profile", plate_no=plate_no, tab="jobs"))


@fleet_bp.route("/fleet/staff")
@_login_required("admin")
def fleet_staff_list():
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    _import_field_staff_from_sqlite(db)
    _import_maintenance_staff_from_sqlite(db)
    _import_orphaned_maintenance_jobs(db)

    unsynced = db.execute("""
        SELECT fs.* FROM field_staff fs
        LEFT JOIN technicians t ON t.technician_code = fs.staff_id
        WHERE t.technician_code IS NULL
    """).fetchall()
    for row in unsynced:
        pw_hash = row["password_hash"] or generate_password_hash("changeme123")
        _sync_field_staff_to_technician(
            db, row["staff_id"], row["full_name"],
            row["phone"] or "", row["username"],
            pw_hash, row["is_active"],
        )
    if unsynced:
        db.commit()

    staff_list = db.execute("""
        SELECT fs.*,
            COALESCE(ec.entry_count, 0) AS entry_count,
            COALESCE(ac.advance_count, 0) AS advance_count
        FROM field_staff fs
        LEFT JOIN (
            SELECT technician_code, COUNT(*) AS entry_count
            FROM maintenance_papers GROUP BY technician_code
        ) ec ON ec.technician_code = fs.staff_id
        LEFT JOIN (
            SELECT staff_code, COUNT(*) AS advance_count
            FROM maintenance_staff_advances GROUP BY staff_code
        ) ac ON ac.staff_code = fs.staff_id
        WHERE fs.staff_id IS NOT NULL AND fs.staff_id != ''
        ORDER BY fs.full_name
    """).fetchall()
    return render_template("fleet/fleet_staff_list.html", staff_list=staff_list)


@fleet_bp.route("/fleet/staff/add", methods=["GET", "POST"])
@_login_required("admin")
def fleet_staff_add():
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()

    if request.method == "POST":
        staff_id = request.form.get("staff_id", "").strip().upper()
        full_name = request.form.get("full_name", "").strip()
        phone = request.form.get("phone", "").strip()
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        photo_file = request.files.get("profile_photo")

        photo_data = None
        photo_content_type = None
        if photo_file and photo_file.filename:
            photo_bytes = photo_file.read()
            if photo_bytes:
                import base64
                photo_data = base64.b64encode(photo_bytes).decode("utf-8")
                photo_content_type = photo_file.content_type or "image/jpeg"

        if not staff_id or not full_name or not username or not password:
            flash("Staff ID, name, username, and password are required.", "error")
            return render_template("fleet/fleet_staff_form.html", page_title="Register New Field Staff", submit_label="Register Staff", s=request.form)

        existing = db.execute("SELECT staff_id FROM field_staff WHERE staff_id = ?", (staff_id,)).fetchone()
        if existing:
            flash("Staff ID already exists.", "error")
            return render_template("fleet/fleet_staff_form.html", page_title="Register New Field Staff", submit_label="Register Staff", s=request.form)

        existing_user = db.execute("SELECT staff_id FROM field_staff WHERE username = ?", (username,)).fetchone()
        if existing_user:
            flash("Username already taken.", "error")
            return render_template("fleet/fleet_staff_form.html", page_title="Register New Field Staff", submit_label="Register Staff", s=request.form)

        pw_hash = generate_password_hash(password)
        db.execute(
            "INSERT INTO field_staff (staff_id, full_name, phone, username, password_hash, photo_data, photo_content_type) VALUES (?,?,?,?,?,?,?)",
            (staff_id, full_name, phone, username, pw_hash, photo_data, photo_content_type),
        )
        _sync_field_staff_to_technician(db, staff_id, full_name, phone, username, pw_hash, 1)
        db.commit()
        flash(f"Staff {full_name} added.", "success")
        return redirect(url_for("fleet.fleet_staff_list"))

    return render_template("fleet/fleet_staff_form.html", page_title="Register New Field Staff", submit_label="Register Staff", s={})


@fleet_bp.route("/fleet/staff/<staff_id>/delete", methods=["POST"])
@_login_required("admin")
def fleet_staff_delete(staff_id):
    _touch_admin_workspace("fleet")
    db = open_db()
    s = db.execute("SELECT staff_id, full_name, phone, username, password_hash, is_active FROM field_staff WHERE staff_id = ?", (staff_id,)).fetchone()
    if not s:
        flash("Staff not found.", "error")
        return redirect(url_for("fleet.fleet_staff_list"))
    # Nullify references so submitted data is preserved
    ALLOWED_TABLES = {"maintenance_jobs", "maintenance_papers", "maintenance_staff_advances"}
    ALLOWED_FIELDS = {"staff_id", "technician_code", "staff_code"}
    for tbl, col in [("maintenance_jobs", "staff_id"), ("maintenance_papers", "technician_code"), ("maintenance_staff_advances", "staff_code")]:
        if tbl not in ALLOWED_TABLES or col not in ALLOWED_FIELDS:
            continue
        try:
            db.execute(f"UPDATE {tbl} SET {col}='' WHERE {col}=?", (staff_id,))
        except Exception:
            try:
                db.execute(f"UPDATE {tbl} SET {col}=NULL WHERE {col}=?", (staff_id,))
            except Exception:
                pass
    db.execute("DELETE FROM field_staff WHERE staff_id = ?", (staff_id,))
    db.execute("DELETE FROM technicians WHERE technician_code = ?", (staff_id,))
    db.commit()
    flash(f"Staff {s['full_name']} deleted. Submitted data preserved.", "success")
    return redirect(url_for("fleet.fleet_staff_list"))


@fleet_bp.route("/fleet/staff/<staff_id>/edit", methods=["GET", "POST"])
@_login_required("admin")
def fleet_staff_edit(staff_id):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    s = db.execute("SELECT staff_id, full_name, phone, username, password_hash, is_active, photo_data, photo_content_type FROM field_staff WHERE staff_id = ?", (staff_id,)).fetchone()
    if not s:
        flash("Staff not found.", "error")
        return redirect(url_for("fleet.fleet_staff_list"))

    if request.method == "POST":
        full_name = request.form.get("full_name", "").strip()
        phone = request.form.get("phone", "").strip()
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        is_active = 1 if request.form.get("is_active") else 0

        photo_data = s["photo_data"]
        photo_content_type = s["photo_content_type"]
        photo_file = request.files.get("profile_photo")
        remove_photo = request.form.get("remove_photo") == "1"
        if remove_photo:
            photo_data = None
            photo_content_type = None
        elif photo_file and photo_file.filename:
            photo_bytes = photo_file.read()
            if photo_bytes:
                import base64
                photo_data = base64.b64encode(photo_bytes).decode("utf-8")
                photo_content_type = photo_file.content_type or "image/jpeg"

        if not full_name or not username:
            flash("Name and username are required.", "error")
            return render_template("fleet/fleet_staff_form.html", page_title="Edit Staff", submit_label="Save Changes", s=request.form)

        if password:
            pw_hash = generate_password_hash(password)
            db.execute("UPDATE field_staff SET full_name=?, phone=?, username=?, password_hash=?, is_active=?, photo_data=?, photo_content_type=? WHERE staff_id=?",
                       (full_name, phone, username, pw_hash, is_active, photo_data, photo_content_type, staff_id))
        else:
            pw_hash = s["password_hash"]
            db.execute("UPDATE field_staff SET full_name=?, phone=?, username=?, is_active=?, photo_data=?, photo_content_type=? WHERE staff_id=?",
                       (full_name, phone, username, is_active, photo_data, photo_content_type, staff_id))
        _sync_field_staff_to_technician(db, staff_id, full_name, phone, username, pw_hash, is_active)
        db.commit()
        flash("Staff updated.", "success")
        return redirect(url_for("fleet.fleet_staff_list"))

    return render_template("fleet/fleet_staff_form.html", page_title="Edit Staff", submit_label="Save Changes", s=s)


# ── ADMIN: Cash Receipts ────────────────────────────────────────


@fleet_bp.route("/fleet/staff/<staff_id>/advances/<int:advance_id>/delete", methods=["POST"])
@_login_required("admin")
def fleet_staff_advance_delete(staff_id, advance_id):
    _touch_admin_workspace("fleet")
    db = open_db()
    a = db.execute("SELECT id, amount, funding_source, reference, notes, entry_date, created_at, staff_code FROM maintenance_staff_advances WHERE id = ? AND staff_code = ?", (advance_id, staff_id)).fetchone()
    if a:
        db.execute("DELETE FROM maintenance_staff_advances WHERE id = ?", (advance_id,))
        db.commit()
        flash("Advance deleted.", "success")
    return redirect(url_for("fleet.fleet_staff_profile", staff_id=staff_id))


@fleet_bp.route("/fleet/staff/<staff_id>/advances/<int:advance_id>/edit", methods=["GET", "POST"])
@_login_required("admin")
def fleet_staff_advance_edit(staff_id, advance_id):
    _touch_admin_workspace("fleet")
    db = open_db()
    s = db.execute("SELECT staff_id, full_name, phone, username, password_hash, is_active FROM field_staff WHERE staff_id = ?", (staff_id,)).fetchone()
    if not s:
        flash("Staff not found.", "error")
        return redirect(url_for("fleet.fleet_staff_list"))
    a = db.execute("SELECT id, amount, funding_source, reference, notes, entry_date, created_at, staff_code FROM maintenance_staff_advances WHERE id = ? AND staff_code = ?", (advance_id, staff_id)).fetchone()
    if not a:
        flash("Advance not found.", "error")
        return redirect(url_for("fleet.fleet_staff_profile", staff_id=staff_id))

    if request.method == "POST":
        amount = request.form.get("amount", "").strip()
        entry_date = request.form.get("entry_date", "").strip()
        entry_time = request.form.get("entry_time", "").strip()
        funding_source = request.form.get("funding_source", "").strip() or "Owner Fund"
        given_by = request.form.get("given_by", "").strip()
        notes = request.form.get("notes", "").strip()

        if not amount:
            flash("Amount is required.", "error")
            return render_template("fleet/fleet_advance_edit.html", s=s, a=a, today=date.today().isoformat(), now=datetime.now())

        full_dt = f"{entry_date} {entry_time}" if entry_time else entry_date
        db.execute(
            "UPDATE maintenance_staff_advances SET amount=?, entry_date=?, funding_source=?, reference=?, notes=? WHERE id=?",
            (float(amount), full_dt, funding_source, given_by or session.get("username", "Admin"), notes or "", advance_id),
        )
        db.commit()
        flash("Advance updated.", "success")
        return redirect(url_for("fleet.fleet_staff_profile", staff_id=staff_id))

    return render_template("fleet/fleet_advance_edit.html", s=s, a=a, today=date.today().isoformat(), now=datetime.now())


# ── ADMIN: Staff Profile ─────────────────────────────────────────

@fleet_bp.route("/fleet/staff/<staff_id>/profile", methods=["GET", "POST"])
@_login_required("admin")
def fleet_staff_profile(staff_id):
    import traceback
    try:
        _touch_admin_workspace("fleet")
        ensure_fleet_tables()
        db = open_db()
        s = db.execute("SELECT staff_id, full_name, phone, username, password_hash, is_active, photo_data, photo_content_type FROM field_staff WHERE staff_id = ?", (staff_id,)).fetchone()
        if not s:
            flash("Staff not found.", "error")
            return redirect(url_for("fleet.fleet_staff_list"))

        if request.method == "POST":
            amount = request.form.get("amount", "").strip()
            entry_date = request.form.get("entry_date", "").strip() or date.today().isoformat()
            entry_time = request.form.get("entry_time", "").strip()
            funding_source = request.form.get("funding_source", "").strip() or "Owner Fund"
            given_by = request.form.get("given_by", "").strip()
            notes = request.form.get("notes", "").strip()

            if not amount:
                flash("Amount is required.", "error")
                return redirect(url_for("fleet.fleet_staff_profile", staff_id=staff_id))

            last = db.execute("SELECT advance_no FROM maintenance_staff_advances ORDER BY id DESC LIMIT 1").fetchone()
            num = 1
            if last:
                num = int(last["advance_no"].split("-")[1]) + 1
            adv_no = f"ADV-{num:04d}"
            full_dt = f"{entry_date} {entry_time}" if entry_time else entry_date
            db.execute(
                "INSERT INTO maintenance_staff_advances (advance_no, staff_code, entry_date, funding_source, amount, reference, notes) VALUES (?,?,?,?,?,?,?)",
                (adv_no, staff_id, full_dt, funding_source, float(amount), given_by or session.get("username", "Admin"), notes or ""),
            )
            db.commit()
            flash(f"AED {amount} given to {s['full_name']}.", "success")
            return redirect(url_for("fleet.fleet_staff_profile", staff_id=staff_id))

        month_filter = request.args.get("month", "")
        date_from = request.args.get("date_from", "")
        date_to = request.args.get("date_to", "")
        vehicle_filter = request.args.get("vehicle", "")
        today_str = date.today().isoformat()
        current_month = today_str[:7]
        filter_month = month_filter[:7] if month_filter else ""

        vehicle_opts = db.execute("""
            SELECT DISTINCT mj.vehicle_id AS plate FROM maintenance_jobs mj WHERE mj.staff_id = ?
            UNION
            SELECT DISTINCT vm.vehicle_no AS plate FROM maintenance_papers mp
            LEFT JOIN vehicle_master vm ON vm.vehicle_id = mp.vehicle_id
            WHERE mp.technician_code = ?
            ORDER BY plate
        """, (staff_id, staff_id)).fetchall()

        card_received = db.execute(
            "SELECT COALESCE(SUM(amount),0) AS t FROM maintenance_staff_advances WHERE staff_code = ?",
            (staff_id,),
        ).fetchone()["t"] or 0

        card_jobs = db.execute(
            "SELECT COALESCE(SUM(COALESCE(staff_amount, amount - tax_amount)),0) AS t FROM maintenance_jobs WHERE staff_id = ? AND status = 'approved'",
            (staff_id,),
        ).fetchone()["t"] or 0

        card_papers = db.execute(
            "SELECT COALESCE(SUM(mp.total_amount),0) AS t FROM maintenance_papers mp WHERE mp.technician_code = ? AND mp.review_status = 'Approved'",
            (staff_id,),
        ).fetchone()["t"] or 0

        card_spent = float(card_jobs) + float(card_papers)
        card_balance = card_received - card_spent

        # Build date-filter WHERE clause
        date_where = ""
        date_params = []
        if filter_month:
            date_where += " AND substr(CAST(mj.created_at AS TEXT),1,7) = ?"
            date_params.append(filter_month)
        else:
            if date_from:
                date_where += " AND substr(CAST(mj.created_at AS TEXT),1,10) >= ?"
                date_params.append(date_from)
            if date_to:
                date_where += " AND substr(CAST(mj.created_at AS TEXT),1,10) <= ?"
                date_params.append(date_to)
        if vehicle_filter:
            date_where += " AND mj.vehicle_id = ?"
            date_params.append(vehicle_filter)
        jobs = db.execute(f"""
            SELECT {_MJ_LIST_COLS}, v.vehicle_type FROM maintenance_jobs mj
            LEFT JOIN vehicles v ON v.plate_no = mj.vehicle_id
            WHERE mj.staff_id = ?{date_where} ORDER BY mj.created_at DESC
        """, (staff_id, *date_params)).fetchall()

        paper_where = ""
        paper_params = []
        if filter_month:
            paper_where += " AND substr(CAST(mp.created_at AS TEXT),1,7) = ?"
            paper_params.append(filter_month)
        else:
            if date_from:
                paper_where += " AND substr(CAST(mp.created_at AS TEXT),1,10) >= ?"
                paper_params.append(date_from)
            if date_to:
                paper_where += " AND substr(CAST(mp.created_at AS TEXT),1,10) <= ?"
                paper_params.append(date_to)
        if vehicle_filter:
            paper_where += " AND vm.vehicle_no = ?"
            paper_params.append(vehicle_filter)
        raw_papers = db.execute(f"""
            SELECT mp.id, mp.paper_no, vm.vehicle_no AS vehicle_id, mp.technician_code AS staff_id,
                   mp.total_amount AS amount, mp.work_summary AS description,
                   mp.review_status AS status, mp.notes AS admin_notes,
                   mp.created_at, 'Maintenance' AS category
            FROM maintenance_papers mp
            LEFT JOIN vehicle_master vm ON vm.vehicle_id = mp.vehicle_id
            WHERE mp.technician_code = ?{paper_where}
            ORDER BY mp.created_at DESC
        """, (staff_id, *paper_params)).fetchall()

        items = []
        for j in jobs:
            d = dict(j)
            d["_type"] = "job"
            d["_id"] = f"job_{j['id']}"
            items.append(d)
        for p in raw_papers:
            d = dict(p)
            if isinstance(d.get("created_at"), datetime):
                d["created_at"] = d["created_at"].strftime("%Y-%m-%d %H:%M:%S")
            d["_type"] = "paper"
            d["_id"] = f"paper_{p['id']}"
            items.append(d)
        items.sort(key=lambda x: str(x.get("created_at", "")), reverse=True)

        adv_where = ""
        adv_params = []
        if filter_month:
            adv_where += " AND substr(entry_date,1,7) = ?"
            adv_params.append(filter_month)
        else:
            if date_from:
                adv_where += " AND substr(entry_date,1,10) >= ?"
                adv_params.append(date_from)
            if date_to:
                adv_where += " AND substr(entry_date,1,10) <= ?"
                adv_params.append(date_to)
        advances = db.execute(f"""
            SELECT id, amount, funding_source, reference, notes, entry_date, created_at, staff_code FROM maintenance_staff_advances WHERE staff_code = ?{adv_where} ORDER BY entry_date DESC
        """, (staff_id, *adv_params)).fetchall()

        cash_items = []
        for a in advances:
            d = dict(a)
            d["_type"] = "advance"
            d["_id"] = f"advance_{a['id']}"
            d["_date"] = str(a.get("entry_date", ""))
            d["_amount"] = a["amount"]
            d["_source"] = a.get("funding_source", "Advance")
            d["_notes"] = a.get("notes", "")
            d["_given_by"] = a.get("reference", "")
            cash_items.append(d)
        cash_items.sort(key=lambda x: x["_date"], reverse=True)

        months = db.execute(
            "SELECT DISTINCT substr(entry_date,1,7) AS m FROM maintenance_staff_advances WHERE staff_code = ? ORDER BY m DESC",
            (staff_id,),
        ).fetchall()

        month_name = "All Time"

        return render_template(
            "fleet/fleet_staff_profile.html",
            s=s,
            card_received=card_received,
            card_spent=card_spent,
            card_balance=card_balance,
            month_name=month_name,
            is_filtered=bool(filter_month),
            items=items,
            cash_items=cash_items,
            months=[r["m"] for r in months],
            filter_month=filter_month,
            date_from=date_from,
            date_to=date_to,
            vehicle_filter=vehicle_filter,
            vehicle_opts=[r["plate"] for r in vehicle_opts if r["plate"]],
            current_month=current_month,
            today=date.today().isoformat(),
            now=datetime.now(),
        )
    except Exception as e:
        tb = traceback.format_exc()
        current_app.logger.error("Profile error for %s:\n%s", staff_id, tb)
        flash(f"Error: {e}", "error")
        return redirect(url_for("fleet.fleet_staff_list"))


@fleet_bp.route("/fleet/staff/<staff_id>/advances/pdf")
def fleet_staff_advances_pdf(staff_id):
    import traceback
    try:
        from app.pdf_service import generate_field_staff_advances_pdf
        _touch_admin_workspace("fleet")
        ensure_fleet_tables()
        db = open_db()
        s = db.execute("SELECT staff_id, full_name, phone, username, password_hash, is_active FROM field_staff WHERE staff_id = ?", (staff_id,)).fetchone()
        if not s:
            flash("Staff not found.", "error")
            return redirect(url_for("fleet.fleet_staff_list"))
        filter_month = request.args.get("month", "")
        date_from = request.args.get("date_from", "")
        date_to = request.args.get("date_to", "")
        date_params = []
        where_adv = ""
        if filter_month:
            where_adv += " AND substr(entry_date,1,7) = ?"
            date_params.append(filter_month[:7])
        else:
            if date_from:
                where_adv += " AND substr(entry_date,1,10) >= ?"
                date_params.append(date_from)
            if date_to:
                where_adv += " AND substr(entry_date,1,10) <= ?"
                date_params.append(date_to)
        if not date_params:
            where_adv = ""
            jp_where = ""
        advances = db.execute(f"SELECT id, amount, funding_source, reference, notes, entry_date, created_at, staff_code FROM maintenance_staff_advances WHERE staff_code = ?{where_adv} ORDER BY entry_date DESC", (staff_id, *date_params)).fetchall()
        total = sum(a["amount"] for a in advances) if advances else 0

        # Jobs/papers filter same period
        jp_where = ""
        jp_params = []
        if filter_month:
            jp_where += " AND substr(CAST(mj.created_at AS TEXT),1,7) = ?"
            jp_params.append(filter_month[:7])
        else:
            if date_from:
                jp_where += " AND substr(CAST(mj.created_at AS TEXT),1,10) >= ?"
                jp_params.append(date_from)
            if date_to:
                jp_where += " AND substr(CAST(mj.created_at AS TEXT),1,10) <= ?"
                jp_params.append(date_to)
        if not jp_params:
            jp_where = ""
        jobs_data = db.execute(f"""
            SELECT mj.id, mj.vehicle_id, mj.amount, mj.created_at, v.plate_no
            FROM maintenance_jobs mj
            LEFT JOIN vehicles v ON v.plate_no = mj.vehicle_id
            WHERE mj.staff_id = ? AND mj.status = 'approved'{jp_where}
            ORDER BY mj.created_at DESC
        """, (staff_id, *jp_params)).fetchall()
        pp_where = jp_where.replace("mj.created_at", "mp.created_at")
        papers_data = db.execute(f"""
            SELECT mp.id, vm.vehicle_no AS vehicle_id, mp.total_amount AS amount, mp.created_at
            FROM maintenance_papers mp
            LEFT JOIN vehicle_master vm ON vm.vehicle_id = mp.vehicle_id
            WHERE mp.technician_code = ? AND mp.review_status = 'Approved'{pp_where}
            ORDER BY mp.created_at DESC
        """, (staff_id, *jp_params)).fetchall()

        from pathlib import Path
        from flask import current_app
        output_dir = Path(current_app.config["GENERATED_DIR"]) / "staff_advances"
        import os
        company = db.execute("SELECT company_name, legal_name, trade_license_no, trade_license_expiry, trn_no, vat_status, address, phone_number, email, bank_name, bank_account_name, bank_account_number, iban, swift_code, invoice_terms, base_currency, logo_data, logo_type, theme_color FROM company_profile LIMIT 1").fetchone()
        base_url = request.host_url.rstrip("/")
        pdf_path = generate_field_staff_advances_pdf(s, advances, jobs_data, papers_data, total, filter_month, date_from, date_to, str(output_dir), current_app.config["STATIC_ASSETS_DIR"], company_profile=company, base_url=base_url)
        relative_path = Path(pdf_path).relative_to(current_app.config["GENERATED_DIR"]).as_posix()
        return redirect(url_for("public_file", filename=relative_path))
    except Exception as e:
        tb = traceback.format_exc()
        current_app.logger.error("Profile PDF error for %s:\n%s", staff_id, tb)
        flash(f"PDF Error: {e}", "error")
        return redirect(url_for("fleet.fleet_staff_profile", staff_id=staff_id))


@fleet_bp.route("/fleet/staff/<staff_id>/jobs/pdf")
def fleet_staff_jobs_pdf(staff_id):
    import traceback
    try:
        from app.pdf_service import generate_field_staff_jobs_pdf
        _touch_admin_workspace("fleet")
        ensure_fleet_tables()
        db = open_db()
        s = db.execute("SELECT staff_id, full_name, phone, username, password_hash, is_active FROM field_staff WHERE staff_id = ?", (staff_id,)).fetchone()
        if not s:
            flash("Staff not found.", "error")
            return redirect(url_for("fleet.fleet_staff_list"))
        filter_month = request.args.get("month", "")
        date_from = request.args.get("date_from", "")
        date_to = request.args.get("date_to", "")
        vehicle_filter = request.args.get("vehicle", "")
        jp_where = ""
        jp_params = []
        if filter_month:
            jp_where += " AND substr(CAST(mj.created_at AS TEXT),1,7) = ?"
            jp_params.append(filter_month[:7])
        else:
            if date_from:
                jp_where += " AND substr(CAST(mj.created_at AS TEXT),1,10) >= ?"
                jp_params.append(date_from)
            if date_to:
                jp_where += " AND substr(CAST(mj.created_at AS TEXT),1,10) <= ?"
                jp_params.append(date_to)
        if vehicle_filter:
            jp_where += " AND mj.vehicle_id = ?"
            jp_params.append(vehicle_filter)
        jobs = db.execute(f"""
            SELECT {_MJ_LIST_COLS}, v.vehicle_type FROM maintenance_jobs mj
            LEFT JOIN vehicles v ON v.plate_no = mj.vehicle_id
            WHERE mj.staff_id = ?{jp_where}
            ORDER BY mj.created_at DESC
        """, (staff_id, *jp_params)).fetchall()
        total_amount = sum(j["amount"] for j in jobs) if jobs else 0
        from pathlib import Path
        from flask import current_app
        output_dir = Path(current_app.config["GENERATED_DIR"]) / "staff_jobs"
        company = db.execute("SELECT company_name, legal_name, trade_license_no, trade_license_expiry, trn_no, vat_status, address, phone_number, email, bank_name, bank_account_name, bank_account_number, iban, swift_code, invoice_terms, base_currency, logo_data, logo_type, theme_color FROM company_profile LIMIT 1").fetchone()
        base_url = request.host_url.rstrip("/")
        pdf_path = generate_field_staff_jobs_pdf(s, jobs, total_amount, filter_month, date_from, date_to, str(output_dir), current_app.config["STATIC_ASSETS_DIR"], company_profile=company, base_url=base_url)
        relative_path = Path(pdf_path).relative_to(current_app.config["GENERATED_DIR"]).as_posix()
        return redirect(url_for("public_file", filename=relative_path))
    except Exception as e:
        tb = traceback.format_exc()
        current_app.logger.error("Jobs PDF error for %s:\n%s", staff_id, tb)
        flash(f"PDF Error: {e}", "error")
        return redirect(url_for("fleet.fleet_staff_profile", staff_id=staff_id))


@fleet_bp.route("/fleet/staff/<staff_id>/delete-data", methods=["POST"])
@_login_required("admin")
def fleet_staff_delete_data(staff_id):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    s = db.execute("SELECT staff_id, full_name, phone, username, password_hash, is_active FROM field_staff WHERE staff_id = ?", (staff_id,)).fetchone()
    if not s:
        flash("Staff not found.", "error")
        return redirect(url_for("fleet.fleet_staff_list"))
    types = request.form.getlist("delete_types")
    deleted = []
    if "advances" in types:
        db.execute("DELETE FROM maintenance_staff_advances WHERE staff_code = ?", (staff_id,))
        deleted.append("Advances")
    if "jobs" in types:
        db.execute("DELETE FROM maintenance_jobs WHERE staff_id = ?", (staff_id,))
        deleted.append("Portal Jobs")
    if "papers" in types:
        paper_nos = [r[0] for r in db.execute("SELECT paper_no FROM maintenance_papers WHERE technician_code = ?", (staff_id,)).fetchall()]
        for pn in paper_nos:
            db.execute("DELETE FROM maintenance_paper_lines WHERE paper_no = ?", (pn,))
        db.execute("DELETE FROM maintenance_papers WHERE technician_code = ?", (staff_id,))
        deleted.append("Maintenance Papers")
    db.commit()
    if deleted:
        flash(f"Deleted: {', '.join(deleted)} for {s['full_name']}.", "success")
    else:
        flash("No data type selected.", "error")
    return redirect(url_for("fleet.fleet_staff_profile", staff_id=staff_id))


@fleet_bp.route("/fleet/staff/<staff_id>/delete-items", methods=["POST"])
@_login_required("admin")
def fleet_staff_delete_items(staff_id):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    s = db.execute("SELECT staff_id, full_name, phone, username, password_hash, is_active FROM field_staff WHERE staff_id = ?", (staff_id,)).fetchone()
    if not s:
        flash("Staff not found.", "error")
        return redirect(url_for("fleet.fleet_staff_list"))
    item_ids = request.form.getlist("item_ids")
    deleted = []
    for item in item_ids:
        parts = item.split("_", 1)
        if len(parts) != 2:
            continue
        prefix, record_id = parts
        try:
            record_id = int(record_id)
        except ValueError:
            continue
        if prefix == "job":
            db.execute("DELETE FROM maintenance_jobs WHERE id = ? AND staff_id = ?", (record_id, staff_id))
            deleted.append(f"job #{record_id}")
        elif prefix == "paper":
            row = db.execute("SELECT paper_no FROM maintenance_papers WHERE id = ? AND technician_code = ?", (record_id, staff_id)).fetchone()
            if row:
                db.execute("DELETE FROM maintenance_paper_lines WHERE paper_no = ?", (row[0],))
            db.execute("DELETE FROM maintenance_papers WHERE id = ? AND technician_code = ?", (record_id, staff_id))
            deleted.append(f"paper #{record_id}")
        elif prefix == "advance":
            db.execute("DELETE FROM maintenance_staff_advances WHERE id = ? AND staff_code = ?", (record_id, staff_id))
            deleted.append(f"advance #{record_id}")
    db.commit()
    if deleted:
        flash(f"Deleted {len(deleted)} item(s): {', '.join(deleted[:10])}{'...' if len(deleted) > 10 else ''}", "success")
    else:
        flash("No items selected.", "error")
    return redirect(url_for("fleet.fleet_staff_profile", staff_id=staff_id))


# ── ADMIN: Pending Approvals ────────────────────────────────────

@fleet_bp.route("/fleet/approvals")
@_login_required("admin")
def fleet_approvals():
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()

    pending_jobs = db.execute(
        f"""SELECT {_MJ_LIST_COLS}, v.vehicle_type, COALESCE(v.plate_no, mj.vehicle_id) AS plate_no, s.full_name AS staff_name, s.staff_id AS staff_key
           FROM maintenance_jobs mj
           LEFT JOIN vehicles v ON v.plate_no = mj.vehicle_id
           JOIN field_staff s ON s.staff_id = mj.staff_id
           WHERE mj.status = 'pending'
           ORDER BY s.full_name ASC, mj.created_at DESC""",
    ).fetchall()

    pending_total = float(
        db.execute("SELECT COALESCE(SUM(amount),0) FROM maintenance_jobs WHERE status = 'pending'").fetchone()[0] or 0
    )

    groups = []
    for j in pending_jobs:
        key = j["staff_key"] or j["staff_id"] or "?"
        name = j["staff_name"] or key
        grp = next((g for g in groups if g["staff_id"] == key), None)
        if grp is None:
            grp = {"staff_id": key, "staff_name": name, "jobs": [], "total": 0.0, "photo_url": None}
            groups.append(grp)
        grp["jobs"].append(j)
        grp["total"] += float(j["amount"] or 0)

    for g in groups:
        try:
            prow = db.execute("SELECT photo_data, photo_content_type FROM field_staff WHERE staff_id = ?", (g["staff_id"],)).fetchone()
            if prow and prow["photo_data"] and prow["photo_content_type"]:
                g["photo_url"] = f"data:{prow['photo_content_type']};base64,{prow['photo_data']}"
        except Exception:
            pass

    recent_approved = db.execute(
        f"""SELECT {_MJ_LIST_COLS}, COALESCE(v.plate_no, mj.vehicle_id) AS plate_no, s.full_name AS staff_name
           FROM maintenance_jobs mj
           LEFT JOIN vehicles v ON v.plate_no = mj.vehicle_id
           JOIN field_staff s ON s.staff_id = mj.staff_id
           WHERE mj.status IN ('approved','rejected')
           ORDER BY mj.created_at DESC LIMIT 20""",
    ).fetchall()

    return render_template("fleet/fleet_approvals.html", pending_jobs=pending_jobs, pending_total=pending_total, recent_approved=recent_approved, pending_groups=groups)


@fleet_bp.route("/fleet/staff-papers-report")
@_login_required("admin")
def fleet_staff_papers_report():
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    month = (request.args.get("month") or "").strip()
    staff_filter = (request.args.get("staff_id") or "").strip()

    where = ["1=1"]
    params = []
    if month:
        where.append("SUBSTR(COALESCE(NULLIF(mj.paper_date, ''), SUBSTR(CAST(mj.created_at AS TEXT), 1, 10)), 1, 7) = ?")
        params.append(month)
    if staff_filter:
        where.append("mj.staff_id = ?")
        params.append(staff_filter)
    clause = " AND ".join(where)

    paper_rows = db.execute(
        f"""
        SELECT mj.id, mj.staff_id, mj.vehicle_id, mj.paper_date, mj.category, mj.description,
               mj.status, mj.supplier_name, mj.supplier_trn, mj.supplier_bill_no,
               mj.tax_mode, mj.tax_amount, mj.amount,
               (mj.amount - mj.tax_amount) AS net_amount,
               COALESCE(v.plate_no, mj.vehicle_id, '-') AS plate_no,
               COALESCE(s.full_name, mj.staff_id, '-') AS staff_name,
               s.photo_data, s.photo_content_type
        FROM maintenance_jobs mj
        LEFT JOIN vehicles v ON v.plate_no = mj.vehicle_id
        LEFT JOIN field_staff s ON s.staff_id = mj.staff_id
        WHERE {clause}
        ORDER BY COALESCE(s.full_name, mj.staff_id) ASC,
                 COALESCE(NULLIF(mj.paper_date, ''), SUBSTR(CAST(mj.created_at AS TEXT), 1, 10)) DESC,
                 mj.id DESC
        """,
        params,
    ).fetchall()

    staff_list = db.execute(
        "SELECT staff_id, full_name, is_active FROM field_staff ORDER BY full_name ASC"
    ).fetchall()
    months = db.execute(
        """
        SELECT DISTINCT SUBSTR(COALESCE(NULLIF(paper_date, ''), SUBSTR(CAST(created_at AS TEXT), 1, 10)), 1, 7) AS m
        FROM maintenance_jobs ORDER BY m DESC
        """
    ).fetchall()

    groups = []
    for r in paper_rows:
        key = r["staff_id"] or "?"
        grp = next((g for g in groups if g["staff_id"] == key), None)
        if grp is None:
            grp = {
                "staff_id": key,
                "staff_name": r["staff_name"] or key,
                "photo_url": None,
                "papers": [],
                "count": 0,
                "net_total": 0.0,
                "vat_total": 0.0,
                "total": 0.0,
                "tax_count": 0,
                "no_tax_count": 0,
            }
            groups.append(grp)
        grp["papers"].append(r)
        grp["count"] += 1
        grp["net_total"] += float(r["net_amount"] or 0)
        grp["vat_total"] += float(r["tax_amount"] or 0)
        grp["total"] += float(r["amount"] or 0)
        if (r["tax_mode"] or "") == "Tax Invoice":
            grp["tax_count"] += 1
        else:
            grp["no_tax_count"] += 1
        if r["photo_data"] and r["photo_content_type"] and not grp["photo_url"]:
            grp["photo_url"] = f"data:{r['photo_content_type']};base64,{r['photo_data']}"

    summary = {
        "count": sum(g["count"] for g in groups),
        "net_total": sum(g["net_total"] for g in groups),
        "vat_total": sum(g["vat_total"] for g in groups),
        "total": sum(g["total"] for g in groups),
        "tax_count": sum(g["tax_count"] for g in groups),
        "no_tax_count": sum(g["no_tax_count"] for g in groups),
    }
    return render_template(
        "fleet/fleet_staff_papers.html",
        groups=groups,
        summary=summary,
        staff_list=staff_list,
        months=[m["m"] for m in months],
        month=month,
        staff_filter=staff_filter,
    )


@fleet_bp.route("/fleet/jobs/approve-all", methods=["POST"])
@_login_required("admin")
def fleet_job_approve_all():
    _touch_admin_workspace("fleet")
    db = open_db()
    pending = db.execute("SELECT id, category, amount, staff_id FROM maintenance_jobs WHERE status='pending'").fetchall()
    if not pending:
        flash("No pending jobs to approve.", "info")
        return redirect(url_for("fleet.fleet_approvals"))
    now = datetime.now().isoformat()
    db.execute("UPDATE maintenance_jobs SET status='approved', approved_at=?, staff_amount=amount-tax_amount WHERE status='pending'", (now,))
    db.commit()
    try:
        from app.notification_service import add_notification
        add_notification(title=f"All {len(pending)} pending jobs approved", type="success", role="admin", link="/fleet/approvals")
        notified = set()
        for j in pending:
            sid = j["staff_id"]
            if sid and sid not in notified:
                notified.add(sid)
                add_notification(title=f"Your job{'s' if len(pending)>1 else ''} approved", type="success", role="technician")
    except Exception:
        pass
    flash(f"All {len(pending)} pending jobs approved.", "success")
    return redirect(url_for("fleet.fleet_approvals"))


@fleet_bp.route("/fleet/jobs/approve-all/<staff_id>", methods=["POST"])
@_login_required("admin")
def fleet_job_approve_all_staff(staff_id):
    _touch_admin_workspace("fleet")
    db = open_db()
    pending = db.execute(
        "SELECT id, category, amount, staff_id FROM maintenance_jobs WHERE status='pending' AND staff_id = ?",
        (staff_id,),
    ).fetchall()
    if not pending:
        flash("No pending jobs for this staff member.", "info")
        return redirect(url_for("fleet.fleet_approvals"))
    now = datetime.now().isoformat()
    db.execute("UPDATE maintenance_jobs SET status='approved', approved_at=?, staff_amount=amount-tax_amount WHERE status='pending' AND staff_id = ?", (now, staff_id))
    db.commit()
    staff_name = "Field Staff"
    try:
        s = db.execute("SELECT full_name FROM field_staff WHERE staff_id = ?", (staff_id,)).fetchone()
        if s:
            staff_name = s[0]
    except Exception:
        pass
    try:
        from app.notification_service import add_notification
        add_notification(title=f"{len(pending)} job(s) for {staff_name} approved", type="success", role="admin", link="/fleet/approvals")
        add_notification(title=f"Your job{'s' if len(pending)>1 else ''} approved", type="success", role="technician")
    except Exception:
        pass
    flash(f"Approved {len(pending)} pending job(s) for {staff_name}.", "success")
    return redirect(url_for("fleet.fleet_approvals"))


@fleet_bp.route("/fleet/jobs/<int:job_id>/approve", methods=["POST"])
@_login_required("admin")
def fleet_job_approve(job_id):
    _touch_admin_workspace("fleet")
    db = open_db()
    job = db.execute("SELECT id, vehicle_id, staff_id, amount, category, description, attachment_name, attachment_data, attachment_type, status, admin_notes, approved_at FROM maintenance_jobs WHERE id = ?", (job_id,)).fetchone()
    if not job:
        flash("Job not found.", "error")
        return redirect(url_for("fleet.fleet_approvals"))
    db.execute(
        "UPDATE maintenance_jobs SET status = 'approved', approved_at = ?, staff_amount = amount - tax_amount WHERE id = ?",
        (datetime.now().isoformat(), job_id),
    )
    db.commit()
    try:
        from app.notification_service import add_notification
        add_notification(
            title=f"Job #{job_id} approved",
            type="success",
            role="admin",
            message=job.get("category","") + " — AED " + str(job.get("amount","")),
            link="/fleet/approvals",
        )
        try:
            staff_id = job.get("staff_id")
            staff_name = "Technician"
            if staff_id:
                s = db.execute("SELECT full_name FROM field_staff WHERE staff_id=?", (staff_id,)).fetchone()
                if s: staff_name = s[0]
        except Exception:
            staff_name = "Technician"
        add_notification(
            title=f"Job #{job_id} approved by admin",
            type="success",
            role="technician",
            message=job.get("category","") + " — AED " + str(job.get("amount","")),
        )
    except Exception:
        pass
    flash(f"Job #{job_id} approved.", "success")
    return redirect(url_for("fleet.fleet_approvals"))


@fleet_bp.route("/fleet/jobs/<int:job_id>/reject", methods=["POST"])
@_login_required("admin")
def fleet_job_reject(job_id):
    _touch_admin_workspace("fleet")
    db = open_db()
    notes = request.form.get("admin_notes", "").strip() or "Rejected by admin"
    db.execute(
        "UPDATE maintenance_jobs SET status = 'rejected', admin_notes = ? WHERE id = ?",
        (notes, job_id),
    )
    db.commit()
    try:
        from app.notification_service import add_notification
        add_notification(
            title=f"Job #{job_id} rejected",
            type="error",
            role="admin",
            message=notes,
            link="/fleet/approvals",
        )
        try:
            job = db.execute("SELECT id, vehicle_id, staff_id, amount, category, description, attachment_name, attachment_data, attachment_type, status, admin_notes, approved_at FROM maintenance_jobs WHERE id = ?", (job_id,)).fetchone()
            staff_id = job.get("staff_id") if job else None
        except Exception:
            staff_id = None
        add_notification(
            title=f"Job #{job_id} rejected by admin",
            type="error",
            role="technician",
            message=notes,
        )
    except Exception:
        pass
    flash(f"Job #{job_id} rejected.", "info")
    return redirect(url_for("fleet.fleet_approvals"))


# ── Serve Attachment ────────────────────────────────────────────

@fleet_bp.route("/fleet/attachment/<int:job_id>")
@_login_required("admin")
def fleet_attachment(job_id):
    db = open_db()
    job = db.execute("SELECT attachment_data, attachment_name, attachment_type FROM maintenance_jobs WHERE id = ?", (job_id,)).fetchone()
    if not job or not job["attachment_data"]:
        flash("Attachment not found.", "error")
        return redirect(url_for("fleet.fleet_approvals"))
    import base64
    from io import BytesIO
    data = base64.b64decode(job["attachment_data"])
    resp = send_file(
        BytesIO(data),
        mimetype=job["attachment_type"] or "application/octet-stream",
        as_attachment=False,
        download_name=job["attachment_name"] or f"attachment_{job_id}",
    )
    resp.headers["Cache-Control"] = "public, max-age=604800"
    return resp


@fleet_bp.route("/fleet/attachment/<int:job_id>/thumb")
@_login_required("admin")
def fleet_attachment_thumb(job_id):
    db = open_db()
    job = db.execute("SELECT attachment_data, attachment_name, attachment_type FROM maintenance_jobs WHERE id = ?", (job_id,)).fetchone()
    if not job or not job["attachment_data"]:
        return ("", 404)
    import base64
    from io import BytesIO
    data = base64.b64decode(job["attachment_data"])
    mime = job["attachment_type"] or "application/octet-stream"
    try:
        from PIL import Image, ImageOps
        img = None
        if mime.startswith("image"):
            img = Image.open(BytesIO(data))
        elif mime == "application/pdf" or (job["attachment_name"] or "").lower().endswith(".pdf"):
            # PDF: render first page so it can show as an inline preview too
            from pdf2image import convert_from_bytes
            pages = convert_from_bytes(data, first_page=1, last_page=1, dpi=60)
            img = pages[0] if pages else None
        if img is not None:
            img = ImageOps.exif_transpose(img)
            img.thumbnail((480, 480))
            if img.mode not in ("RGB", "L"):
                img = img.convert("RGB")
            out = BytesIO()
            img.save(out, format="JPEG", quality=72)
            data = out.getvalue()
            mime = "image/jpeg"
    except Exception:
        pass
    resp = send_file(
        BytesIO(data),
        mimetype=mime,
        as_attachment=False,
        download_name=f"thumb_{job_id}",
    )
    resp.headers["Cache-Control"] = "public, max-age=31536000, immutable"
    resp.headers["ETag"] = f'"j{job_id}-{len(data)}"'
    return resp


# ═════════════════════════════════════════════════════════════════
# FUEL MANAGEMENT
# ═════════════════════════════════════════════════════════════════

@fleet_bp.route("/fleet/fuel")
@_login_required("admin")
def fuel_list():
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    vehicle_filter = request.args.get("vehicle", "")
    month_filter = request.args.get("month", "")
    params = []
    where = ""
    if vehicle_filter:
        where += " AND fe.vehicle_plate = ?"
        params.append(vehicle_filter)
    if month_filter:
        where += " AND substr(fe.entry_date,1,7) = ?"
        params.append(month_filter)
    entries = db.execute(f"""
        SELECT fe.* FROM fuel_entries fe
        WHERE 1=1{where}
        ORDER BY fe.entry_date DESC, fe.id DESC
    """, params).fetchall()
    vehicles = db.execute("SELECT plate_no FROM vehicles ORDER BY plate_no").fetchall()
    total_gallons = sum(e["gallons"] for e in entries) if entries else 0
    total_amount = sum(e["total_amount"] for e in entries) if entries else 0
    return render_template(
        "fleet/fuel_list.html",
        entries=entries,
        vehicles=vehicles,
        vehicle_filter=vehicle_filter,
        month_filter=month_filter,
        total_gallons=total_gallons,
        total_amount=total_amount,
    )


@fleet_bp.route("/fleet/fuel/report/pdf")
@_login_required("admin")
def fuel_report_pdf():
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    vehicle_filter = request.args.get("vehicle", "")
    month_filter = request.args.get("month", "")
    params = []
    where = ""
    if vehicle_filter:
        where += " AND fe.vehicle_plate = ?"
        params.append(vehicle_filter)
    if month_filter:
        where += " AND substr(fe.entry_date,1,7) = ?"
        params.append(month_filter)
    entries = db.execute(f"""
        SELECT fe.* FROM fuel_entries fe
        WHERE 1=1{where}
        ORDER BY fe.entry_date DESC, fe.id DESC
    """, params).fetchall()
    output_dir = current_app.config.get("GENERATED_BACKUP_DIR", "/tmp")
    company = db.execute("SELECT company_name, legal_name, trade_license_no, trade_license_expiry, trn_no, vat_status, address, phone_number, email, bank_name, bank_account_name, bank_account_number, iban, swift_code, invoice_terms, base_currency, logo_data, logo_type, theme_color FROM company_profile LIMIT 1").fetchone()
    cp = dict(company) if company else {}
    assets_dir = str(Path(current_app.root_path).parent / "app" / "static")
    pdf_path = generate_fuel_report_pdf(
        entries=entries,
        vehicle_filter=vehicle_filter,
        month_filter=month_filter,
        output_dir=output_dir,
        assets_dir=assets_dir,
        company_profile=cp,
    )
    return send_file(pdf_path, as_attachment=True, download_name=f"fuel_report_{month_filter or 'all'}.pdf")


@fleet_bp.route("/fleet/fuel/add", methods=["GET", "POST"])
@_login_required("admin")
def fuel_add():
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    if request.method == "POST":
        supplier_id = request.form.get("supplier_id", "").strip()
        if not supplier_id:
            flash("Please select a supplier.", "error")
            return redirect(url_for("fleet.fuel_add"))
        supplier = db.execute("SELECT id, supplier_name FROM suppliers WHERE id = ?", (supplier_id,)).fetchone()
        if not supplier:
            flash("Supplier not found.", "error")
            return redirect(url_for("fleet.fuel_add"))

        # ── Bulk mode: multiple vehicle rows (vehicle_plate[]) ──
        bulk_plates = request.form.getlist("vehicle_plate[]")
        if bulk_plates and bulk_plates[0].strip():
            entry_dates = request.form.getlist("entry_date[]")
            gallons_list = request.form.getlist("gallons[]")
            amounts_list = request.form.getlist("total_amount[]")
            notes_list = request.form.getlist("notes[]")
            global_notes = request.form.get("notes", "").strip()
            added = 0
            for i, raw_plate in enumerate(bulk_plates):
                plate = raw_plate.strip()
                edate = entry_dates[i].strip() if i < len(entry_dates) else ""
                try:
                    gln = float(gallons_list[i] or 0) if i < len(gallons_list) else 0
                except (TypeError, ValueError):
                    gln = 0
                amt_raw = amounts_list[i].strip() if i < len(amounts_list) else ""
                try:
                    amt = float(amt_raw or 0)
                except (TypeError, ValueError):
                    amt = 0
                note = (notes_list[i].strip() if i < len(notes_list) else "") or global_notes
                if not plate or not edate or gln <= 0 or amt <= 0:
                    continue
                rate = round(amt / gln, 3)
                _insert_fuel_entry(db, plate, edate, gln, rate, supplier, note)
                added += 1
            db.commit()
            if added:
                flash(f"{added} fuel entr{'y' if added == 1 else 'ies'} added.", "success")
            else:
                flash("No valid rows to add. Each row needs a vehicle, date, gallons and amount.", "error")
            return redirect(url_for("fleet.fuel_list"))

        # ── Single mode (used by the vehicle-profile quick form) ──
        vehicle_plate = request.form.get("vehicle_plate", "").strip()
        entry_date = request.form.get("entry_date", "").strip()
        try:
            gallons = float(request.form.get("gallons", 0) or 0)
        except (TypeError, ValueError):
            gallons = 0
        try:
            rate = float(request.form.get("rate_per_gallon", 0) or 0)
        except (TypeError, ValueError):
            rate = 0
        # Auto-calculate rate from total_amount if provided
        total_amount = request.form.get("total_amount", "").strip()
        if total_amount:
            try:
                amt = float(total_amount or 0)
            except (TypeError, ValueError):
                amt = 0
            if rate <= 0 and gallons > 0:
                rate = round(amt / gallons, 3)
        notes = request.form.get("notes", "").strip()
        if not vehicle_plate or not entry_date or gallons <= 0 or rate <= 0:
            flash("Please fill all required fields.", "error")
            return redirect(url_for("fleet.fuel_add"))
        fuel_id, total = _insert_fuel_entry(db, vehicle_plate, entry_date, gallons, rate, supplier, notes)
        db.commit()
        flash(f"Fuel entry added: {gallons} GLN × {rate} = AED {total}", "success")
        return redirect(url_for("fleet.fuel_list"))
    vehicles = db.execute("SELECT plate_no FROM vehicles ORDER BY plate_no").fetchall()
    suppliers = db.execute("SELECT id, supplier_name FROM suppliers WHERE status = 'Active' ORDER BY supplier_name").fetchall()
    today = date.today()
    return render_template(
        "fleet/fuel_form.html",
        vehicles=vehicles,
        suppliers=suppliers,
        entry=None,
        vehicles_json=json.dumps([v["plate_no"] for v in vehicles]),
        today=today.isoformat(),
        today_month=today.strftime("%Y-%m"),
    )


def _insert_fuel_entry(db, vehicle_plate, entry_date, gallons, rate, supplier, notes):
    """Insert one fuel entry plus its linked supplier_expense; returns (fuel_id, total)."""
    total = round(gallons * rate, 2)
    fuel_desc = f"{gallons} GLN × {rate} = AED {total} — {vehicle_plate}"
    if db.backend == "postgres":
        fuel_result = db.execute(
            """INSERT INTO fuel_entries (vehicle_plate, entry_date, gallons, rate_per_gallon, total_amount, supplier_id, supplier_name, notes)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?) RETURNING id""",
            (vehicle_plate, entry_date, gallons, rate, total, supplier["id"], supplier["supplier_name"], notes),
        )
        fuel_id = fuel_result.fetchone()[0]
        exp_result = db.execute(
            """INSERT INTO supplier_expenses (supplier_id, expense_date, amount, category, description, earning_type, quantity, rate, vehicle_no, status)
               VALUES (?, ?, ?, 'Fuel', ?, 'Fuel', ?, ?, ?, 'approved') RETURNING id""",
            (supplier["id"], entry_date, total, notes or fuel_desc, gallons, rate, vehicle_plate),
        )
        expense_id = exp_result.fetchone()[0]
    else:
        db.execute(
            """INSERT INTO fuel_entries (vehicle_plate, entry_date, gallons, rate_per_gallon, total_amount, supplier_id, supplier_name, notes)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (vehicle_plate, entry_date, gallons, rate, total, supplier["id"], supplier["supplier_name"], notes),
        )
        fuel_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
        db.execute(
            """INSERT INTO supplier_expenses (supplier_id, expense_date, amount, category, description, earning_type, quantity, rate, vehicle_no, status)
               VALUES (?, ?, ?, 'Fuel', ?, 'Fuel', ?, ?, ?, 'approved')""",
            (supplier["id"], entry_date, total, notes or fuel_desc, gallons, rate, vehicle_plate),
        )
        expense_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
    db.execute("UPDATE fuel_entries SET source_expense_id = ? WHERE id = ?", (expense_id, fuel_id))
    return fuel_id, total


@fleet_bp.route("/fleet/fuel/<int:entry_id>/edit", methods=["GET", "POST"])
@_login_required("admin")
def fuel_edit(entry_id):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    entry = db.execute("SELECT id, vehicle_plate, entry_date, gallons, rate_per_gallon, total_amount, supplier_id, supplier_name, notes, source_expense_id, created_at FROM fuel_entries WHERE id = ?", (entry_id,)).fetchone()
    if not entry:
        flash("Fuel entry not found.", "error")
        return redirect(url_for("fleet.fuel_list"))
    if request.method == "POST":
        vehicle_plate = request.form.get("vehicle_plate", "").strip()
        entry_date = request.form.get("entry_date", "").strip()
        gallons = float(request.form.get("gallons", 0) or 0)
        rate = float(request.form.get("rate_per_gallon", 0) or 0)
        # Auto-calculate rate from total_amount if provided
        total_amount = request.form.get("total_amount", "").strip()
        if total_amount:
            amt = float(total_amount or 0)
            if rate <= 0 and gallons > 0:
                rate = round(amt / gallons, 3)
        supplier_id = request.form.get("supplier_id", "").strip()
        notes = request.form.get("notes", "").strip()
        if not vehicle_plate or not entry_date or gallons <= 0 or rate <= 0 or not supplier_id:
            flash("Please fill all required fields.", "error")
            return redirect(url_for("fleet.fuel_edit", entry_id=entry_id))
        supplier = db.execute("SELECT id, supplier_name FROM suppliers WHERE id = ?", (supplier_id,)).fetchone()
        if not supplier:
            flash("Supplier not found.", "error")
            return redirect(url_for("fleet.fuel_edit", entry_id=entry_id))
        total = round(gallons * rate, 2)
        db.execute(
            """UPDATE fuel_entries SET vehicle_plate=?, entry_date=?, gallons=?, rate_per_gallon=?, total_amount=?, supplier_id=?, supplier_name=?, notes=?
               WHERE id=?""",
            (vehicle_plate, entry_date, gallons, rate, total, supplier["id"], supplier["supplier_name"], notes, entry_id),
        )
        # Update linked supplier_expenses
        if entry["source_expense_id"]:
            fuel_desc = f"{gallons} GLN × {rate} = AED {total} — {vehicle_plate}"
            db.execute(
                """UPDATE supplier_expenses SET expense_date=?, amount=?, quantity=?, rate=?, vehicle_no=?, description=?, earning_type=?, category=?
                   WHERE id=?""",
                (entry_date, total, gallons, rate, vehicle_plate, notes or fuel_desc, 'Fuel', 'Fuel', entry["source_expense_id"]),
            )
        db.commit()
        flash("Fuel entry updated.", "success")
        return redirect(url_for("fleet.fuel_list"))
    vehicles = db.execute("SELECT plate_no FROM vehicles ORDER BY plate_no").fetchall()
    suppliers = db.execute("SELECT id, supplier_name FROM suppliers WHERE status = 'Active' ORDER BY supplier_name").fetchall()
    return render_template("fleet/fuel_form.html", vehicles=vehicles, suppliers=suppliers, entry=entry)


@fleet_bp.route("/fleet/fuel/<int:entry_id>/delete", methods=["POST"])
@_login_required("admin")
def fuel_delete(entry_id):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    entry = db.execute("SELECT id, vehicle_plate, entry_date, gallons, rate_per_gallon, total_amount, supplier_id, supplier_name, notes, source_expense_id, created_at FROM fuel_entries WHERE id = ?", (entry_id,)).fetchone()
    if not entry:
        flash("Fuel entry not found.", "error")
        return redirect(url_for("fleet.fuel_list"))
    if entry["source_expense_id"]:
        db.execute("DELETE FROM supplier_expenses WHERE id = ?", (entry["source_expense_id"],))
    db.execute("DELETE FROM fuel_entries WHERE id = ?", (entry_id,))
    db.commit()
    flash("Fuel entry deleted.", "info")
    return redirect(url_for("fleet.fuel_list"))


@fleet_bp.route("/fleet/fuel/supplier/<int:supplier_id>")
@_login_required("admin")
def fuel_supplier_statement(supplier_id):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    supplier = db.execute("SELECT id, supplier_name FROM suppliers WHERE id = ?", (supplier_id,)).fetchone()
    if not supplier:
        flash("Supplier not found.", "error")
        return redirect(url_for("fleet.fuel_list"))
    month_filter = request.args.get("month", "")
    date_from = request.args.get("date_from", "")
    date_to = request.args.get("date_to", "")
    params = [supplier_id]
    where = ""
    if month_filter:
        where += " AND substr(fe.entry_date,1,7) = ?"
        params.append(month_filter[:7])
    else:
        if date_from:
            where += " AND fe.entry_date >= ?"
            params.append(date_from)
        if date_to:
            where += " AND fe.entry_date <= ?"
            params.append(date_to)
    entries = db.execute(f"""
        SELECT fe.* FROM fuel_entries fe
        WHERE fe.supplier_id = ?{where}
        ORDER BY fe.entry_date DESC, fe.id DESC
    """, params).fetchall()
    total_gallons = sum(e["gallons"] for e in entries) if entries else 0
    total_amount = sum(e["total_amount"] for e in entries) if entries else 0
    return render_template(
        "fleet/fuel_supplier_statement.html",
        supplier=supplier,
        entries=entries,
        month_filter=month_filter,
        date_from=date_from,
        date_to=date_to,
        total_gallons=total_gallons,
        total_amount=total_amount,
    )


@fleet_bp.route("/fleet/vat-quick", methods=["GET", "POST"])
@_login_required("admin")
def vat_quick():
    _touch_admin_workspace("fleet")
    db = open_db()
    vehicles = db.execute("SELECT plate_no, vehicle_type FROM vehicles ORDER BY vehicle_type, plate_no").fetchall()

    selected = (request.args.get("vehicle") or "").strip()
    sort_field = (request.args.get("field") or "date").lower()
    if sort_field not in ("date", "amount"):
        sort_field = "date"
    sort_dir = (request.args.get("sort") or "desc").lower()
    if sort_dir not in ("asc", "desc"):
        sort_dir = "desc"
    order_col = "mj.created_at" if sort_field == "date" else "mj.amount"
    sort_clause = f"ORDER BY {order_col} ASC, mj.id ASC" if sort_dir == "asc" else f"ORDER BY {order_col} DESC, mj.id DESC"
    results = []
    summary = {"total": 0.0, "vat": 0.0, "without_tax": 0, "with_vat": 0}
    if selected:
        results = db.execute(
            f"""SELECT {_MJ_LIST_COLS}, (mj.amount - mj.tax_amount) AS net_amount,
                       COALESCE(s.full_name, 'Admin') AS staff_name
                FROM maintenance_jobs mj
                LEFT JOIN field_staff s ON s.staff_id = mj.staff_id
                WHERE mj.vehicle_id = ? AND mj.status IN ('approved','pending','rejected')
                {sort_clause}""",
            (selected,),
        ).fetchall()
        results = [dict(r) for r in results]
        summary["total"] = round(sum(float(r["amount"] or 0) for r in results), 2)
        summary["vat"] = round(sum(float(r["tax_amount"] or 0) for r in results), 2)
        summary["with_vat"] = sum(1 for r in results if r["tax_mode"] == "Tax Invoice")
        summary["without_tax"] = sum(1 for r in results if r["tax_mode"] != "Tax Invoice")

    if request.method == "POST":
        action = request.form.get("action")
        if action == "add_new_job":
            try:
                net = round(float(request.form.get("new_amount") or 0), 2)
            except ValueError:
                net = 0.0
            category = (request.form.get("new_category") or "Other").strip()
            description = (request.form.get("new_description") or "").strip()
            entry_date = (request.form.get("new_date") or "").strip() or date.today().isoformat()
            tax_mode = (request.form.get("new_tax_mode") or "Without Tax").strip()
            supplier_name = (request.form.get("new_supplier_name") or "").strip()
            supplier_bill_no = (request.form.get("new_supplier_bill_no") or "").strip()
            supplier_trn = (request.form.get("new_supplier_trn") or "").strip()
            if net <= 0:
                flash("Amount must be greater than zero.", "error")
                return redirect(url_for("fleet.vat_quick", vehicle=selected or None))
            if tax_mode == "Tax Invoice":
                tax_amount = round(net * 0.05, 2)
                amount_total = round(net + tax_amount, 2)
            else:
                tax_amount = 0.0
                amount_total = net
            db.execute(
                "INSERT INTO maintenance_jobs (vehicle_id, staff_id, amount, category, description, status, tax_mode, tax_amount, staff_amount, created_at, supplier_name, supplier_trn, supplier_bill_no) "
                "VALUES (?, 'admin', ?, ?, ?, 'approved', ?, ?, ?, ?, ?, ?, ?)",
                (selected, amount_total, category, description, tax_mode, tax_amount, net, entry_date, supplier_name or None, supplier_trn or None, supplier_bill_no or None),
            )
            if supplier_name:
                _upsert_maintenance_supplier(db, supplier_name, supplier_trn or None)
            db.commit()
            flash(f"New job added for {selected} (net {net:.2f}{' + VAT ' + format(tax_amount, '.2f') if tax_mode == 'Tax Invoice' else ''}).", "success")
            return redirect(url_for("fleet.vat_quick", vehicle=selected or None))

        job_id = request.form.get("job_id")
        try:
            job_id = int(job_id) if job_id else None
        except (TypeError, ValueError):
            job_id = None
        job = db.execute("SELECT * FROM maintenance_jobs WHERE id = ?", (job_id,)).fetchone() if job_id else None
        if job is None:
            flash("Job not found.", "error")
            return redirect(url_for("fleet.vat_quick", vehicle=selected or None))
        if action == "add_vat":
            try:
                net = round(float(request.form.get("net_amount") or 0), 2)
            except ValueError:
                net = round(float(job["amount"]) - float(job["tax_amount"] or 0), 2)
            if net < 0:
                net = 0.0
            tax = round(net * 0.05, 2)
            total = round(net + tax, 2)
            supplier_name = (request.form.get("supplier_name") or "").strip() or job["supplier_name"]
            supplier_bill_no = (request.form.get("supplier_bill_no") or "").strip() or job["supplier_bill_no"]
            supplier_trn = (request.form.get("supplier_trn") or "").strip() or job["supplier_trn"]
            db.execute(
                """UPDATE maintenance_jobs
                   SET amount=?, tax_mode='Tax Invoice', tax_amount=?,
                       supplier_name=?, supplier_bill_no=?, supplier_trn=?
                   WHERE id=?""",
                (total, tax, supplier_name or None, supplier_bill_no or None, supplier_trn or None, job_id),
            )
            if supplier_name:
                _upsert_maintenance_supplier(db, supplier_name, supplier_trn or None)
            flash(f"Job #{job_id}: VAT 5% added (net {net:.2f} + VAT {tax:.2f} = {total:.2f}). Staff balance untouched.", "success")
        elif action == "remove_vat":
            net = round(float(job["amount"]) - float(job["tax_amount"] or 0), 2)
            db.execute(
                "UPDATE maintenance_jobs SET amount=?, tax_mode='Without Tax', tax_amount=0 WHERE id=?",
                (net, job_id),
            )
            flash(f"Job #{job_id}: VAT removed (back to {net:.2f}). Staff balance untouched.", "success")
        db.commit()
        return redirect(url_for("fleet.vat_quick", vehicle=selected or None))

    _ensure_maintenance_suppliers_table(db)
    supplier_suggestions = []
    try:
        supplier_rows = db.execute(
            "SELECT DISTINCT supplier_name FROM maintenance_jobs WHERE supplier_name IS NOT NULL AND supplier_name != '' "
            "UNION SELECT DISTINCT supplier_name FROM maintenance_papers WHERE supplier_name IS NOT NULL AND supplier_name != '' "
            "UNION SELECT DISTINCT name FROM maintenance_suppliers WHERE name IS NOT NULL AND name != '' "
            "ORDER BY supplier_name ASC"
        ).fetchall()
        supplier_suggestions = [r[0] for r in supplier_rows]
    except Exception:
        try:
            supplier_rows = db.execute(
                "SELECT DISTINCT supplier_name FROM maintenance_jobs WHERE supplier_name IS NOT NULL AND supplier_name != '' "
                "UNION SELECT DISTINCT name FROM maintenance_suppliers WHERE name IS NOT NULL AND name != '' "
                "ORDER BY supplier_name ASC"
            ).fetchall()
            supplier_suggestions = [r[0] for r in supplier_rows]
        except Exception:
            supplier_suggestions = []

    return render_template(
        "fleet/vat_quick.html",
        vehicles=vehicles,
        selected=selected,
        sort_dir=sort_dir,
        sort_field=sort_field,
        results=results,
        summary=summary,
        supplier_suggestions=supplier_suggestions,
    )


@fleet_bp.route("/fleet/vat-pending", methods=["GET", "POST"])
@_login_required("admin")
def vat_pending():
    _touch_admin_workspace("fleet")
    db = open_db()

    if request.method == "POST":
        action = request.form.get("action")
        job_id = request.form.get("job_id")
        try:
            job_id = int(job_id) if job_id else None
        except (TypeError, ValueError):
            job_id = None
        job = db.execute("SELECT * FROM maintenance_jobs WHERE id = ?", (job_id,)).fetchone() if job_id else None
        if job is None:
            flash("Job not found.", "error")
            return redirect(url_for("fleet.vat_pending"))
        if action == "add_vat":
            try:
                net = round(float(request.form.get("net_amount") or 0), 2)
            except ValueError:
                net = round(float(job["amount"]) - float(job["tax_amount"] or 0), 2)
            if net < 0:
                net = 0.0
            tax = round(net * 0.05, 2)
            total = round(net + tax, 2)
            supplier_name = (request.form.get("supplier_name") or "").strip() or job["supplier_name"]
            supplier_bill_no = (request.form.get("supplier_bill_no") or "").strip() or job["supplier_bill_no"]
            supplier_trn = (request.form.get("supplier_trn") or "").strip() or job["supplier_trn"]
            job_date = (request.form.get("job_date") or "").strip()
            if job_date:
                ts = str(job["created_at"] or "")[:19]
                if len(ts) >= 11 and (ts[10] == " " or "T" in ts):
                    new_created_at = job_date + ts[10:19]
                else:
                    new_created_at = job_date
                db.execute(
                    """UPDATE maintenance_jobs
                       SET amount=?, tax_mode='Tax Invoice', tax_amount=?,
                           supplier_name=?, supplier_bill_no=?, supplier_trn=?,
                           created_at=?, paper_date=?, vat_check=NULL
                       WHERE id=?""",
                    (total, tax, supplier_name or None, supplier_bill_no or None, supplier_trn or None,
                     new_created_at, job_date, job_id),
                )
            else:
                db.execute(
                    """UPDATE maintenance_jobs
                       SET amount=?, tax_mode='Tax Invoice', tax_amount=?,
                           supplier_name=?, supplier_bill_no=?, supplier_trn=?, vat_check=NULL
                       WHERE id=?""",
                    (total, tax, supplier_name or None, supplier_bill_no or None, supplier_trn or None, job_id),
                )
            if supplier_name:
                _upsert_maintenance_supplier(db, supplier_name, supplier_trn or None)
            flash(f"Job #{job_id}: VAT 5% added (net {net:.2f} + VAT {tax:.2f} = {total:.2f}). Staff balance untouched — removed from pending list.", "success")
        elif action == "no_vat":
            db.execute("UPDATE maintenance_jobs SET vat_check='no_vat' WHERE id=?", (job_id,))
            flash(f"Job #{job_id} marked as No VAT — hidden from pending list (nothing changed).", "success")
        elif action == "unhide":
            db.execute("UPDATE maintenance_jobs SET vat_check=NULL WHERE id=?", (job_id,))
            flash(f"Job #{job_id} restored to pending list.", "success")
        db.commit()
        return redirect(url_for("fleet.vat_pending",
                                show_hidden=request.form.get("show_hidden") or "",
                                month=request.form.get("month") or ""))

    month = (request.args.get("month") or "").strip()
    show_hidden = (request.args.get("show_hidden") or "").strip()
    hide_clause = "AND (mj.vat_check IS NULL OR mj.vat_check != 'no_vat')"
    if show_hidden:
        hide_clause = ""
    rows = db.execute(
        f"""SELECT {_MJ_LIST_COLS}, (mj.amount - mj.tax_amount) AS net_amount,
                   COALESCE(s.full_name, 'Admin') AS staff_name
            FROM maintenance_jobs mj
            LEFT JOIN field_staff s ON s.staff_id = mj.staff_id
            WHERE mj.status = 'approved'
              AND (mj.tax_mode IS NULL OR mj.tax_mode != 'Tax Invoice')
              AND mj.attachment_data IS NOT NULL AND mj.attachment_data != ''
              {hide_clause}
            ORDER BY mj.created_at DESC, mj.id DESC""",
    ).fetchall()
    rows = [dict(r) for r in rows]
    for r in rows:
        r["job_date"] = str(r.get("created_at") or "")[:10]
    if month:
        rows = [r for r in rows if (str(r.get("created_at") or "")[:7]) == month]

    month_options = []
    try:
        d = datetime.now()
        for i in range(24):
            ym = (d.year, d.month)
            month_options.append(f"{ym[0]:04d}-{ym[1]:02d}")
            if d.month == 1:
                d = d.replace(year=d.year - 1, month=12)
            else:
                d = d.replace(month=d.month - 1)
    except Exception:
        pass

    summary = {
        "total": round(sum(float(r["amount"] or 0) for r in rows), 2),
        "count": len(rows),
        "with_photo": sum(1 for r in rows if r["has_attachment"]),
        "hidden": 0,
        "month": month,
    }
    if not show_hidden:
        hidden_rows = db.execute(
            "SELECT created_at FROM maintenance_jobs WHERE status='approved' AND (tax_mode IS NULL OR tax_mode != 'Tax Invoice') AND vat_check='no_vat' AND attachment_data IS NOT NULL AND attachment_data != ''"
        ).fetchall()
        if month:
            hidden_rows = [r for r in hidden_rows if (str(r[0] or "")[:7]) == month]
        summary["hidden"] = len(hidden_rows)

    _ensure_maintenance_suppliers_table(db)
    supplier_suggestions = []
    try:
        supplier_rows = db.execute(
            "SELECT DISTINCT supplier_name FROM maintenance_jobs WHERE supplier_name IS NOT NULL AND supplier_name != '' "
            "UNION SELECT DISTINCT supplier_name FROM maintenance_papers WHERE supplier_name IS NOT NULL AND supplier_name != '' "
            "UNION SELECT DISTINCT name FROM maintenance_suppliers WHERE name IS NOT NULL AND name != '' "
            "ORDER BY supplier_name ASC"
        ).fetchall()
        supplier_suggestions = [r[0] for r in supplier_rows]
    except Exception:
        supplier_suggestions = []

    return render_template(
        "fleet/vat_pending.html",
        rows=rows,
        summary=summary,
        supplier_suggestions=supplier_suggestions,
        show_hidden=show_hidden,
        month=month,
        month_options=month_options,
    )


@fleet_bp.route("/fleet/jobs/<int:job_id>/edit", methods=["GET", "POST"])
@_login_required("admin")
def fleet_job_edit(job_id):
    _touch_admin_workspace("fleet")
    db = open_db()
    job = db.execute("""SELECT mj.*, v.vehicle_type, fs.full_name as staff_name
                        FROM maintenance_jobs mj
                        LEFT JOIN vehicles v ON v.plate_no = mj.vehicle_id
                        JOIN field_staff fs ON fs.staff_id = mj.staff_id
                        WHERE mj.id = ?""", (job_id,)).fetchone()
    if not job:
        flash("Job not found.", "error")
        return redirect(url_for("fleet.fleet_approvals"))

    vehicles = db.execute("SELECT plate_no, vehicle_type, model, year, ownership_type, partner_name, partner_percent, status, notes FROM vehicles ORDER BY vehicle_type, plate_no").fetchall()
    _ensure_maintenance_suppliers_table(db)
    supplier_rows = db.execute(
        "SELECT DISTINCT supplier_name FROM maintenance_jobs WHERE supplier_name IS NOT NULL AND supplier_name != '' UNION SELECT DISTINCT supplier_name FROM maintenance_papers WHERE supplier_name IS NOT NULL AND supplier_name != '' UNION SELECT DISTINCT name FROM maintenance_suppliers WHERE name IS NOT NULL AND name != '' ORDER BY supplier_name ASC"
    ).fetchall()
    supplier_suggestions = [r[0] for r in supplier_rows]
    supplier_trn_rows = db.execute(
        "SELECT supplier_name, supplier_trn FROM maintenance_jobs WHERE supplier_name IS NOT NULL AND supplier_name != '' AND supplier_trn IS NOT NULL AND supplier_trn != '' UNION SELECT supplier_name, supplier_trn FROM maintenance_papers WHERE supplier_name IS NOT NULL AND supplier_name != '' AND supplier_trn IS NOT NULL AND supplier_trn != '' UNION SELECT name, trn FROM maintenance_suppliers WHERE trn IS NOT NULL AND trn != ''"
    ).fetchall()
    supplier_trn_map = {}
    for row in supplier_trn_rows:
        if row[0] and not supplier_trn_map.get(row[0]):
            supplier_trn_map[row[0]] = row[1]

    if request.method == "POST":
        vehicle_id = request.form.get("vehicle_id", "").strip()
        amount = request.form.get("amount", "").strip()
        category = request.form.get("category", "").strip()
        description = request.form.get("description", "").strip()
        status = request.form.get("status", "").strip()
        supplier_name = request.form.get("supplier_name", "").strip()
        supplier_trn = request.form.get("supplier_trn", "").strip()
        supplier_bill_no = request.form.get("supplier_bill_no", "").strip()
        tax_mode = request.form.get("tax_mode", job["tax_mode"] or "Without Tax").strip() or "Without Tax"
        job_date = request.form.get("job_date", "").strip()

        if not amount or not category:
            flash("Amount and category are required.", "error")
            return render_template("fleet/fleet_job_edit.html", job=job, vehicles=vehicles, categories=MAINTENANCE_CATEGORIES, supplier_suggestions=supplier_suggestions, supplier_trn_map=supplier_trn_map)
        if tax_mode == "Tax Invoice":
            if not supplier_name:
                flash("Workshop name is required when the bill includes VAT 5%.", "error")
                return render_template("fleet/fleet_job_edit.html", job=job, vehicles=vehicles, categories=MAINTENANCE_CATEGORIES, supplier_suggestions=supplier_suggestions, supplier_trn_map=supplier_trn_map)
            if not supplier_bill_no:
                flash("Bill number is required when the bill includes VAT 5%.", "error")
                return render_template("fleet/fleet_job_edit.html", job=job, vehicles=vehicles, categories=MAINTENANCE_CATEGORIES, supplier_suggestions=supplier_suggestions, supplier_trn_map=supplier_trn_map)
        try:
            net_amount = round(float(amount), 2)
        except ValueError:
            net_amount = 0.0
        if tax_mode == "Tax Invoice":
            tax_amount = round(net_amount * 0.05, 2)
            amount_total = round(net_amount + tax_amount, 2)
        else:
            tax_amount = 0.0
            amount_total = net_amount

        attachment_name = job["attachment_name"]
        attachment_data = job["attachment_data"]
        attachment_type = job["attachment_type"]
        if "attachment" in request.files:
            file = request.files["attachment"]
            if file.filename:
                import base64
                attachment_name = file.filename
                attachment_data = base64.b64encode(file.read()).decode("utf-8")
                attachment_type = file.content_type

        new_status = status if status in ("pending", "approved", "rejected") else job["status"]
        staff_amount = job["staff_amount"]
        if new_status == "approved" and staff_amount is None:
            staff_amount = round((job["amount"] or 0) - (job["tax_amount"] or 0), 2)

        new_created_at = job["created_at"]
        new_paper_date = job["paper_date"]
        if job_date:
            ts = str(job["created_at"] or "")[:19]
            if len(ts) >= 11 and (ts[10] == " " or "T" in ts):
                new_created_at = job_date + ts[10:19]
            else:
                new_created_at = job_date
            new_paper_date = job_date

        db.execute(
            """UPDATE maintenance_jobs
               SET vehicle_id=?, amount=?, category=?, description=?,
                   attachment_name=?, attachment_data=?, attachment_type=?,
                   supplier_name=?, supplier_trn=?, supplier_bill_no=?,
                   tax_mode=?, tax_amount=?, staff_amount=?,
                   status=?, created_at=?, paper_date=?
               WHERE id=?""",
            (vehicle_id or "N/A", amount_total, category, description,
             attachment_name, attachment_data, attachment_type,
             supplier_name or None, supplier_trn or None, supplier_bill_no or None,
             tax_mode, tax_amount, staff_amount, new_status, new_created_at, new_paper_date, job_id),
        )
        db.commit()
        if supplier_name:
            _upsert_maintenance_supplier(db, supplier_name, supplier_trn)
        flash("Job updated.", "success")
        return redirect(url_for("fleet.fleet_approvals"))

    return render_template(
        "fleet/fleet_job_edit.html",
        job=job,
        vehicles=vehicles,
        categories=MAINTENANCE_CATEGORIES,
        supplier_suggestions=supplier_suggestions,
        supplier_trn_map=supplier_trn_map,
        job_date=str(job["created_at"] or "")[:10],
    )


@fleet_bp.route("/fleet/jobs/<int:job_id>/delete", methods=["POST"])
@_login_required("admin")
def fleet_job_delete(job_id):
    _touch_admin_workspace("fleet")
    db = open_db()
    job = db.execute("SELECT id FROM maintenance_jobs WHERE id = ?", (job_id,)).fetchone()
    if not job:
        flash("Job not found.", "error")
    else:
        db.execute("DELETE FROM maintenance_jobs WHERE id = ?", (job_id,))
        db.commit()
        flash("Job deleted.", "info")
    return redirect(url_for("fleet.fleet_approvals"))


@fleet_bp.route("/fleet/jobs/<int:job_id>/revert", methods=["POST"])
@_login_required("admin")
def fleet_job_revert(job_id):
    _touch_admin_workspace("fleet")
    db = open_db()
    db.execute("UPDATE maintenance_jobs SET status='pending' WHERE id=?", (job_id,))
    db.commit()
    flash("Job reverted to pending.", "success")
    return redirect(url_for("fleet.fleet_approvals"))


# ═════════════════════════════════════════════════════════════════
# ENDPOINT: Maintenance Suppliers registry + profile
# ═════════════════════════════════════════════════════════════════

@fleet_bp.route("/fleet/maintenance-suppliers")
@_login_required("admin")
def fleet_maintenance_suppliers():
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    _ensure_maintenance_suppliers_table(db)

    reg_rows = db.execute("SELECT name, trn FROM maintenance_suppliers").fetchall()
    reg = {r["name"]: r["trn"] for r in reg_rows}

    paper_agg = db.execute(
        """
        SELECT COALESCE(NULLIF(p.supplier_name, ''), '') AS name,
               COUNT(*) AS cnt,
               SUM(CASE WHEN p.tax_mode = 'Tax Invoice' THEN 1 ELSE 0 END) AS tax_cnt,
               SUM(p.subtotal) AS net,
               SUM(p.tax_amount) AS vat,
               SUM(p.total_amount) AS gross,
               MAX(COALESCE(NULLIF(p.paper_date, ''), SUBSTR(CAST(p.created_at AS TEXT), 1, 10), '')) AS last_date,
               MAX(p.supplier_trn) AS trn
        FROM maintenance_papers p
        WHERE p.review_status = 'Approved' AND COALESCE(NULLIF(p.supplier_name, ''), '') != ''
        GROUP BY 1
        """
    ).fetchall()
    job_agg = db.execute(
        """
        SELECT COALESCE(NULLIF(mj.supplier_name, ''), '') AS name,
               COUNT(*) AS cnt,
               SUM(CASE WHEN mj.tax_mode = 'Tax Invoice' THEN 1 ELSE 0 END) AS tax_cnt,
               SUM(mj.amount - mj.tax_amount) AS net,
               SUM(mj.tax_amount) AS vat,
               SUM(mj.amount) AS gross,
               MAX(COALESCE(NULLIF(mj.paper_date, ''), SUBSTR(CAST(mj.created_at AS TEXT), 1, 10), '')) AS last_date,
               MAX(mj.supplier_trn) AS trn
        FROM maintenance_jobs mj
        WHERE mj.status = 'approved' AND COALESCE(NULLIF(mj.supplier_name, ''), '') != ''
        GROUP BY 1
        """
    ).fetchall()

    def _blank_s():
        return {"cnt": 0, "tax_cnt": 0, "net": 0.0, "vat": 0.0, "gross": 0.0, "last_date": "", "trn": ""}

    stats = {}
    all_names = set(reg)

    def _merge(name, rec):
        s = stats.setdefault(name, _blank_s())
        s["cnt"] += int(rec[0] or 0)
        s["tax_cnt"] += int(rec[1] or 0)
        s["net"] += float(rec[2] or 0)
        s["vat"] += float(rec[3] or 0)
        s["gross"] += float(rec[4] or 0)
        if rec[5] and rec[5] > s["last_date"]:
            s["last_date"] = rec[5]
        s["trn"] = (s["trn"] or reg.get(name)) or (rec[6] or "")

    for r in paper_agg:
        all_names.add(r["name"])
        _merge(r["name"], (r["cnt"], r["tax_cnt"], r["net"], r["vat"], r["gross"], r["last_date"], r["trn"]))
    for r in job_agg:
        all_names.add(r["name"])
        _merge(r["name"], (r["cnt"], r["tax_cnt"], r["net"], r["vat"], r["gross"], r["last_date"], r["trn"]))

    suppliers = []
    for name in sorted(all_names, key=lambda n: n.lower()):
        s = stats.get(name, _blank_s())
        s["name"] = name
        s["trn"] = s["trn"] or reg.get(name) or ""
        suppliers.append(s)
    suppliers.sort(key=lambda s: (s["gross"], s["cnt"]), reverse=True)

    totals = {
        "count": sum(s["cnt"] for s in suppliers),
        "tax_count": sum(s["tax_cnt"] for s in suppliers),
        "net": sum(s["net"] for s in suppliers),
        "vat": sum(s["vat"] for s in suppliers),
        "gross": sum(s["gross"] for s in suppliers),
    }
    return render_template(
        "fleet/fleet_maintenance_suppliers.html",
        suppliers=suppliers,
        totals=totals,
    )


@fleet_bp.route("/fleet/maintenance-supplier/<path:supplier_name>")
@_login_required("admin")
def fleet_maintenance_supplier_profile(supplier_name):
    _touch_admin_workspace("fleet")
    ensure_fleet_tables()
    db = open_db()
    _ensure_maintenance_suppliers_table(db)

    reg = db.execute("SELECT name, trn FROM maintenance_suppliers WHERE name = ?", (supplier_name,)).fetchone()

    papers = db.execute(
        """
        SELECT p.id, p.paper_no, p.paper_date, p.vehicle_no, p.vehicle_id, p.work_summary,
               p.supplier_name, p.supplier_trn, p.supplier_bill_no,
               p.tax_mode, p.subtotal, p.tax_amount, p.total_amount, p.review_status,
               COALESCE(NULLIF(p.paper_date, ''), SUBSTR(CAST(p.created_at AS TEXT), 1, 10), '') AS eff_date
        FROM maintenance_papers p
        WHERE p.supplier_name = ?
        ORDER BY COALESCE(NULLIF(p.paper_date, ''), SUBSTR(CAST(p.created_at AS TEXT), 1, 10), '') DESC, p.id DESC
        """,
        (supplier_name,),
    ).fetchall()
    jobs = db.execute(
        """
        SELECT mj.id, mj.vehicle_id, mj.description, mj.supplier_name, mj.supplier_trn, mj.supplier_bill_no,
               mj.tax_mode, mj.tax_amount, mj.amount, mj.status,
               COALESCE(NULLIF(mj.paper_date, ''), SUBSTR(CAST(mj.created_at AS TEXT), 1, 10), '') AS eff_date
        FROM maintenance_jobs mj
        WHERE mj.supplier_name = ?
        ORDER BY COALESCE(NULLIF(mj.paper_date, ''), SUBSTR(CAST(mj.created_at AS TEXT), 1, 10), '') DESC, mj.id DESC
        """,
        (supplier_name,),
    ).fetchall()

    records = []
    for p in papers:
        records.append({
            "type": "Paper",
            "ref": p["paper_no"] or str(p["id"]),
            "date": p["eff_date"] or "",
            "vehicle": p["vehicle_no"] or p["vehicle_id"] or "",
            "work": p["work_summary"] or "",
            "bill_no": p["supplier_bill_no"] or "",
            "net": float(p["subtotal"] or 0),
            "vat": float(p["tax_amount"] or 0),
            "gross": float(p["total_amount"] or 0),
            "status": p["review_status"] or "Pending",
            "is_tax": (p["tax_mode"] or "") == "Tax Invoice",
        })
    for j in jobs:
        records.append({
            "type": "Job",
            "ref": "JOB-" + str(j["id"]),
            "date": j["eff_date"] or "",
            "vehicle": j["vehicle_id"] or "",
            "work": j["description"] or "",
            "bill_no": j["supplier_bill_no"] or "",
            "net": float((j["amount"] or 0) - (j["tax_amount"] or 0)),
            "vat": float(j["tax_amount"] or 0),
            "gross": float(j["amount"] or 0),
            "status": (j["status"] or "pending").title(),
            "is_tax": (j["tax_mode"] or "") == "Tax Invoice",
        })
    records.sort(key=lambda r: r["date"] or "", reverse=True)

    trn = (reg and reg["trn"]) or ""
    if not trn:
        trn = next((p_s["supplier_trn"] for p_s in papers if p_s["supplier_trn"]), "") or \
              next((j_s["supplier_trn"] for j_s in jobs if j_s["supplier_trn"]), "")
    approved = [r for r in records if r["status"].lower() == "approved"]
    summary = {
        "trn": trn or None,
        "registered": bool(reg),
        "total_records": len(records),
        "approved_count": len(approved),
        "tax_count": sum(1 for r in approved if r["is_tax"]),
        "no_tax_count": sum(1 for r in approved if not r["is_tax"]),
        "net": sum(r["net"] for r in approved),
        "vat": sum(r["vat"] for r in approved),
        "gross": sum(r["gross"] for r in approved),
    }
    return render_template(
        "fleet/fleet_maintenance_supplier_profile.html",
        supplier_name=supplier_name,
        records=records,
        summary=summary,
    )
